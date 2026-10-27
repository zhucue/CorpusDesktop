"""P6-2: corpus source / license metadata, attribution in exports, import and index in rounds."""

import csv
import io

from openpyxl import load_workbook

from tests.api.conftest import drain, upload

TEXT = "Language is a system.\nThe language of science."
META = {
    "license": "CC BY 2.0 FR",
    "attribution": "Tatoeba contributors",
    "source_url": "https://tatoeba.org/",
    "version": "2026-10-10",
}
NOTICE = (
    "来源：Tatoeba；许可：CC BY 2.0 FR；署名：Tatoeba contributors；"
    "链接：https://tatoeba.org/；版本：2026-10-10"
)


def test_corpus_meta_and_export_notice(owner, project, make_user):
    admin = make_user("admin@example.edu", admin=True)
    r = admin.post(
        "/api/corpora", json={"name": "Tatoeba", "language": "zh", "meta": META}
    )
    assert r.status_code == 200, r.text
    builtin = r.json()
    assert builtin["meta"] == META
    assert owner.get(f"/api/corpora/{builtin['id']}").json()["meta"] == META
    too_long = {"name": "x", "language": "en", "meta": {"license": "x" * 201}}
    assert admin.post("/api/corpora", json=too_long).status_code == 422
    plain = owner.post(
        f"/api/projects/{project['id']}/corpora",
        json={"name": "作文", "language": "en"},
    ).json()
    assert plain["meta"] == {}

    upload(admin, builtin["id"], {"a.txt": TEXT})
    upload(owner, plain["id"], {"a.txt": TEXT})
    drain()

    def export(corpus_id, fmt):
        job = owner.post(
            "/api/search/export",
            json={
                "scope": {"corpus_id": corpus_id},
                "query": "language",
                "format": fmt,
                "project_id": project["id"],
            },
        ).json()
        drain()
        return owner.get(f"/api/jobs/{job['id']}/download").content

    rows = list(
        csv.reader(io.StringIO(export(builtin["id"], "csv").decode("utf-8-sig")))
    )
    assert rows[0] == [NOTICE] and rows[1] == [] and rows[2][0] == "命中 ID"
    assert len(rows) == 3 + 2
    sheet = load_workbook(io.BytesIO(export(builtin["id"], "xlsx"))).active
    values = [
        [c for c in row if c is not None] for row in sheet.iter_rows(values_only=True)
    ]
    assert values[0] == [NOTICE] and values[1] == [] and values[2][0] == "命中 ID"

    # corpora without license / attribution export the table only
    rows = list(csv.reader(io.StringIO(export(plain["id"], "csv").decode("utf-8-sig"))))
    assert rows[0][0] == "命中 ID"


def test_import_and_index_in_rounds(owner, project, monkeypatch, fake_engine):
    """Small WRITE_SENTENCES / LOAD_TOKENS force several INSERT and load rounds; results are unchanged."""
    from services.api.storage import corpus_key
    from services.worker import tasks

    monkeypatch.setattr(tasks, "WRITE_SENTENCES", 1)
    monkeypatch.setattr(tasks, "LOAD_TOKENS", 5)
    corpus = owner.post(
        f"/api/projects/{project['id']}/corpora", json={"name": "x", "language": "en"}
    ).json()
    upload(
        owner, corpus["id"], {"a.txt": TEXT, "b.txt": "One two.\nThree four five six."}
    )
    drain()
    c = owner.get(f"/api/corpora/{corpus['id']}").json()
    assert c["status"] == "ready"
    assert c["stats"] == {"documents": 2, "sentences": 4, "tokens": 14}
    indexed = fake_engine.indexes[corpus_key(corpus["id"])]
    by_title = {d.title: d for _, d in indexed}
    assert sorted(by_title) == ["a", "b"]
    a = by_title["a"]
    assert [a.text[s.start : s.end] for s in a.sentences] == [
        "Language is a system.",
        "The language of science.",
    ]
    assert [t.form for t in a.sentences[1].tokens] == [
        "The",
        "language",
        "of",
        "science.",
    ]


def test_load_batches():
    from services.worker import tasks

    rows = [("a", 3), ("b", 2), ("c", 1_500_000), ("d", 1), ("e", 1)]
    assert list(tasks._load_batches(rows)) == [["a", "b"], ["c"], ["d", "e"]]
