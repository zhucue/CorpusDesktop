"""P5 span / relation / alignment layers, multi-annotator agreement, span queries and exports (contracts/p5.md).

The chat model is a fake that marks known learner errors (offsets deliberately off for one of them), relates the
spans of a sentence and links tokens one to one; embeddings are fake topic vectors.
"""

import io
import json
import math
import re
import zipfile
from collections import Counter

import pytest
from lxml import etree
from openpyxl import load_workbook
from sqlalchemy import text

from tests.api.conftest import drain, upload
from tests.api.test_annotations import bearer

ESSAY = (
    "Yesterday I go to school with my friend .\n"
    "She have many book in her bag .\n"
    "We was happy .\n"
    "This sentence is BROKEN for the model .\n"
    "Everything here is fine ."
)
ERRORS = {"go": "时态", "have": "主谓一致", "book": "用词", "was": "主谓一致"}
TMX = [
    ("I like apples .", "我 喜欢 苹果 。"),
    ("The cat sleeps .", "猫 在 睡觉 。"),
    ("Good night .", "晚安 。"),
]
TOPICS = {"apple": 0, "苹果": 0, "cat": 1, "猫": 1, "night": 2, "晚安": 2}


@pytest.fixture(autouse=True)
def fake_embed(monkeypatch):
    from engine.kb import embed

    def vector(t):
        v = [0.0] * 384
        for word, axis in TOPICS.items():
            v[axis] += t.lower().count(word)
        v[5] = 0.05
        n = math.sqrt(sum(x * x for x in v))
        return [x / n for x in v]

    monkeypatch.setattr(
        embed, "embed", lambda texts, download=True: [vector(t) for t in texts]
    )
    monkeypatch.setattr(embed, "load", lambda download=True: None)
    monkeypatch.setattr(embed, "model_name", lambda: "fake-embed")


@pytest.fixture
def fake_llm(monkeypatch):
    from services.worker import annot_tasks, llm

    state = {"calls": [], "prompts": []}

    def spans_answer(user):
        sentences = json.loads(user.split("句子：\n", 1)[1])
        if any("BROKEN" in s["text"] for s in sentences):
            return "抱歉，我无法输出"  # not JSON: retried once, then the batch fails
        out = []
        for s in sentences:
            for m in re.finditer(r"\S+", s["text"]):
                if m.group() in ERRORS:
                    start = m.start() + (3 if m.group() == "have" else 0)  # off by 3
                    out.append(
                        {
                            "id": s["id"],
                            "text": m.group(),
                            "start": start,
                            "end": start + len(m.group()),
                            "label": ERRORS[m.group()],
                            "attrs": {},
                        }
                    )
        out.append({"id": 99, "text": "x", "label": "时态"})  # dropped
        return json.dumps(out, ensure_ascii=False)

    def relations_answer(user):
        ids = [int(n) for n in re.findall(r"句子 (\d+)：", user)]
        return json.dumps(
            [{"id": n, "from": 1, "to": 2, "label": "cause"} for n in ids]
        )

    def words_answer(user):
        pairs = json.loads(user[user.index("\n[") + 1 :])
        out = []
        for p in pairs:
            n = min(len(p["src"].split()), len(p["tgt"].split()))
            out += [{"id": p["id"], "src": [k], "tgt": [k]} for k in range(1, n + 1)]
        out.append({"id": 1, "src": [99], "tgt": [1]})  # out of range: dropped
        return json.dumps(out)

    def chat(system, user, temperature=0.2):
        state["calls"].append(system)
        state["prompts"].append(user)
        if system == annot_tasks.SPAN_SYSTEM:
            return spans_answer(user)
        if system == annot_tasks.RELATION_SYSTEM:
            return relations_answer(user)
        return words_answer(user)

    monkeypatch.setattr(llm, "chat", chat)
    monkeypatch.setattr(llm, "model_id", lambda: "fake-chat")
    return state


def label_set(client, name):
    return next(
        s for s in client.get("/api/label-sets").json()["items"] if s["name"] == name
    )


def sql(query, **params):
    from services.api.db import transaction

    with transaction() as db:
        return db.execute(text(query), params).all()


@pytest.fixture
def essay(owner, project):
    c = owner.post(
        f"/api/projects/{project['id']}/corpora",
        json={"name": "作文", "language": "en"},
    ).json()
    upload(owner, c["id"], {"essay.txt": ESSAY})
    drain()
    c["document_id"] = owner.get(f"/api/corpora/{c['id']}/documents").json()["items"][
        0
    ]["id"]
    return c


def guideline(owner, project, body):
    r = owner.post(
        f"/api/projects/{project['id']}/kb/documents",
        json={"kind": "guideline", "title": "偏误标注规范", "body": body},
    )
    assert r.status_code == 200, r.text
    drain()
    return r.json()


def make_layer(client, project, **body):
    r = client.post(f"/api/projects/{project['id']}/layers", json=body)
    assert r.status_code == 200, r.text
    drain()
    return client.get(f"/api/layers/{r.json()['id']}").json()


def span_layer(owner, project, essay, **extra):
    return make_layer(
        owner,
        project,
        corpus_id=essay["id"],
        kind="span",
        name="偏误",
        label_set_id=label_set(owner, "learner-errors")["id"],
        **extra,
    )


def sentences(client, layer, **params):
    r = client.get(f"/api/layers/{layer['id']}/spans", params=params)
    assert r.status_code == 200, r.text
    return r.json()


def all_spans(client, layer, **params):
    return [
        s for item in sentences(client, layer, **params)["items"] for s in item["spans"]
    ]


def by_text(spans):
    return {s["text"]: s for s in spans}


def search(client, corpus, q, **kw):
    r = client.post(
        "/api/search", json={"scope": {"corpus_id": corpus["id"]}, "query": q, **kw}
    )
    assert r.status_code == 200, r.text
    return r.json()


def effective_count(layer_id, labels, attrs=None):
    """Independent SQL count of effective spans (single mode) with one of the labels."""
    extra = "".join(f" AND attrs->>'{k}' = '{v}'" for k, v in (attrs or {}).items())
    return sql(
        "SELECT count(*) FROM span_annotations WHERE layer_id = CAST(:l AS uuid)"
        " AND status IN ('confirmed', 'modified') AND label = ANY(:labels)" + extra,
        l=layer_id,
        labels=list(labels),
    )[0][0]


def patch_span(client, layer, span, **body):
    r = client.patch(f"/api/layers/{layer['id']}/spans/{span['id']}", json=body)
    assert r.status_code == 200, r.text
    return r.json()


# ---- span layers ----


def test_span_preannotation(owner, project, essay, fake_llm):
    long_rule = "时态：看时间状语。" + "细则" * 1500
    g = guideline(owner, project, long_rule)
    layer = span_layer(owner, project, essay, guideline_document_id=g["id"])
    assert layer["status"] == "ready" and layer["kind"] == "span"
    assert layer["source_model"] == "fake-chat" and layer["mode"] == "single"
    assert layer["annotators"] == [owner.user_id]
    assert layer["guideline_document_id"] == g["id"]
    # one request per batch of ≤ 10 sentences, retried once for the unparseable answer
    assert len(fake_llm["calls"]) == 2
    prompt = fake_llm["prompts"][0]
    assert (
        "时态：看时间状语" in prompt and "主谓一致" in prompt and "上级：偏误" in prompt
    )
    excerpt = prompt.split("标注规范摘录：\n", 1)[1].split("\n请标注", 1)[0].strip()
    assert len(excerpt) == 2000
    failed = layer["progress"]["failed_sentences"]
    assert len(failed) == 5  # the whole batch (one request for five sentences)
    assert layer["progress"]["total"] == 0

    # a smaller batch keeps the other sentences: redo sentences 0–2 only
    r = owner.post(
        f"/api/layers/{layer['id']}/preannotate",
        json={"sentence_range": {"start": 0, "end": 3}},
    )
    assert r.status_code == 200, r.text
    drain()
    job = owner.get(f"/api/jobs/{r.json()['id']}").json()
    assert job["state"] == "succeeded" and job["result"]["written"] == 4
    assert job["result"]["dropped"] == 1
    layer = owner.get(f"/api/layers/{layer['id']}").json()
    assert layer["progress"]["total"] == 4 and layer["progress"]["suggested"] == 4
    assert layer["progress"]["failed_sentences"] == []
    spans = by_text(all_spans(owner, layer))
    assert set(spans) == set(ERRORS)
    have = spans["have"]
    assert (
        ESSAY[have["start_cp"] : have["end_cp"]] == "have"
    )  # wrong offset repaired from the text
    assert have["label"] == "主谓一致" and have["model_label"] == "主谓一致"
    assert have["source"] == "model:fake-chat" and have["status"] == "suggested"
    assert spans["go"]["sentence_index"] == 0 and spans["was"]["sentence_index"] == 2

    # re-run without range: sentences with rows are not sent again (only the failed ones)
    calls = len(fake_llm["calls"])
    r = owner.post(f"/api/layers/{layer['id']}/preannotate", json={})
    drain()
    sent = fake_llm["prompts"][calls:]
    assert sent and all("Yesterday" not in p for p in sent)
    assert owner.get(f"/api/layers/{layer['id']}").json()["progress"]["total"] == 4

    # a range on a token layer is refused
    pos = make_layer(
        owner,
        project,
        corpus_id=essay["id"],
        kind="pos",
        name="pos",
        label_set_id=label_set(owner, "upos")["id"],
    )
    r = owner.post(
        f"/api/layers/{pos['id']}/preannotate",
        json={"sentence_range": {"start": 0, "end": 1}},
    )
    assert r.status_code == 422


def test_layer_validation(owner, project, essay, member):
    pid = project["id"]
    upos = label_set(owner, "upos")["id"]
    errors = label_set(owner, "learner-errors")["id"]
    bad = [
        {"corpus_id": essay["id"], "kind": "span", "name": "x", "label_set_id": upos},
        {"corpus_id": essay["id"], "kind": "span", "name": "x"},
        {
            "kind": "relation",
            "name": "x",
            "label_set_id": label_set(owner, "error-correction")["id"],
        },
        {
            "corpus_id": essay["id"],
            "kind": "alignment",
            "name": "x",
            "label_set_id": errors,
        },
        {
            "corpus_id": essay["id"],
            "kind": "span",
            "name": "x",
            "label_set_id": errors,
            "mode": "multi",
            "annotators": [member("viewer").user_id],
        },
    ]
    for body in bad:
        r = owner.post(f"/api/projects/{pid}/layers", json=body)
        assert r.status_code == 422, (body, r.text)
    note = owner.post(
        f"/api/projects/{pid}/kb/documents",
        json={"kind": "note", "title": "n", "body": "x"},
    ).json()
    r = owner.post(
        f"/api/projects/{pid}/layers",
        json={
            "corpus_id": essay["id"],
            "kind": "span",
            "name": "x",
            "label_set_id": errors,
            "guideline_document_id": note["id"],
        },
    )
    assert r.json()["error"]["code"] == "INVALID_GUIDELINE"
    layer = make_layer(
        owner,
        project,
        corpus_id=essay["id"],
        kind="span",
        name="手工",
        label_set_id=errors,
        preannotate=False,
    )
    assert (
        layer["status"] == "ready"
        and layer["source_model"] == "human"
        and not layer["job_id"]
    )


def test_span_review_and_relations(owner, project, essay, fake_llm, member):
    layer = span_layer(owner, project, essay)
    page = sentences(owner, layer, page_size=2)
    assert [s["sentence_index"] for s in page["items"]] == [0, 1]
    first = page["items"][0]
    assert first["text"] == "Yesterday I go to school with my friend ."
    assert (first["start_cp"], first["end_cp"]) == (0, len(first["text"]))
    rest = sentences(owner, layer, cursor=page["next_cursor"])
    assert [s["sentence_index"] for s in rest["items"]] == [2, 3, 4]
    assert (
        rest["next_cursor"] is None and rest["items"][1]["spans"] == []
    )  # every sentence listed
    pending = sentences(owner, layer, status="pending")["items"]
    assert [
        s["sentence_index"] for s in pending
    ] == []  # the batch failed: nothing suggested

    owner.post(
        f"/api/layers/{layer['id']}/preannotate",
        json={"sentence_range": {"start": 0, "end": 3}},
    )
    drain()
    assert [
        s["sentence_index"] for s in sentences(owner, layer, status="pending")["items"]
    ] == [0, 1, 2]
    spans = by_text(all_spans(owner, layer))
    confirmed = patch_span(owner, layer, spans["go"], action="confirm")
    assert confirmed["status"] == "confirmed" and confirmed["revision"] == 1
    assert confirmed["source"] == f"human:{owner.user_id}"
    patch_span(owner, layer, spans["have"], action="confirm")
    book = patch_span(
        owner, layer, spans["book"], action="modify", attrs={"correction": "books"}
    )
    assert book["status"] == "modified" and book["attrs"] == {"correction": "books"}
    assert (
        patch_span(owner, layer, spans["was"], action="reject")["status"] == "rejected"
    )
    start = ESSAY.index("fine")
    r = owner.post(
        f"/api/layers/{layer['id']}/spans",
        json={
            "document_id": essay["document_id"],
            "start_cp": start,
            "end_cp": start + 4,
            "label": "时态",
            "attrs": {"tense": "past"},
        },
    )
    assert r.status_code == 200, r.text
    fine = r.json()
    assert (
        fine["status"] == "confirmed"
        and fine["text"] == "fine"
        and fine["sentence_index"] == 4
    )
    dup = owner.post(
        f"/api/layers/{layer['id']}/spans",
        json={
            "document_id": essay["document_id"],
            "start_cp": start,
            "end_cp": start + 4,
            "label": "时态",
        },
    )
    assert dup.status_code == 409
    bad = owner.post(
        f"/api/layers/{layer['id']}/spans",
        json={
            "document_id": essay["document_id"],
            "start_cp": 3,
            "end_cp": 2000,
            "label": "时态",
        },
    )
    assert bad.json()["error"]["code"] == "INVALID_RANGE"
    wrong = owner.post(
        f"/api/layers/{layer['id']}/spans",
        json={
            "document_id": essay["document_id"],
            "start_cp": 0,
            "end_cp": 2,
            "label": "NOUN",
        },
    )
    assert wrong.json()["error"]["code"] == "INVALID_LABEL"
    # extent change: a span over two sentences has no sentence index and shows in both
    moved = patch_span(
        owner, layer, fine, action="modify", start_cp=start - 26, end_cp=start + 4
    )
    assert moved["sentence_index"] is None
    moved = patch_span(
        owner, layer, fine, action="modify", start_cp=start, end_cp=start + 4
    )
    assert moved["sentence_index"] == 4 and moved["revision"] == 3
    assert (
        owner.delete(f"/api/layers/{layer['id']}/spans/{spans['go']['id']}").status_code
        == 409
    )
    extra = owner.post(
        f"/api/layers/{layer['id']}/spans",
        json={
            "document_id": essay["document_id"],
            "start_cp": 0,
            "end_cp": 9,
            "label": "用词",
        },
    ).json()
    assert (
        owner.delete(f"/api/layers/{layer['id']}/spans/{extra['id']}").status_code
        == 204
    )

    progress = owner.get(f"/api/layers/{layer['id']}").json()["progress"]
    assert progress == {
        "total": 5,
        "suggested": 0,
        "confirmed": 2,
        "modified": 2,  # book, and fine after the range changes
        "rejected": 1,
        "failed_sentences": [],
    }
    stats = owner.get(f"/api/layers/{layer['id']}/stats").json()
    assert {x["label"]: x["count"] for x in stats["labels"]} == {
        "时态": 2,
        "主谓一致": 1,
        "用词": 1,
    }
    docs = owner.get(f"/api/layers/{layer['id']}/documents").json()["items"]
    assert docs[0]["sentences"] == 5 and docs[0]["pending"] == 0

    viewer = member("viewer")
    assert viewer.get(f"/api/layers/{layer['id']}/spans").status_code == 200
    assert (
        viewer.patch(
            f"/api/layers/{layer['id']}/spans/{fine['id']}", json={"action": "reject"}
        ).status_code
        == 403
    )

    # relation layer over the span layer: the model relates the two spans of sentence 1
    rel = make_layer(
        owner,
        project,
        kind="relation",
        name="更正",
        base_layer_id=layer["id"],
        label_set_id=label_set(owner, "error-correction")["id"],
    )
    assert rel["status"] == "ready" and rel["corpus_id"] == essay["id"]
    items = sentences(owner, rel)["items"]
    assert [len(s["relations"]) for s in items] == [0, 1, 0, 0, 0]
    suggested = items[1]["relations"][0]
    span_ids = {s["id"]: s["text"] for s in items[1]["spans"]}
    assert (span_ids[suggested["from_span_id"]], span_ids[suggested["to_span_id"]]) == (
        "have",
        "book",
    )
    assert suggested["status"] == "suggested" and suggested["label"] == "cause"
    # the review workspace lists the relation layer's documents with pending sentences
    docs = owner.get(f"/api/layers/{rel['id']}/documents")
    assert docs.status_code == 200 and docs.json()["items"][0]["pending"] == 1
    r = owner.patch(
        f"/api/layers/{rel['id']}/relations/{suggested['id']}",
        json={"action": "modify", "label": "correction"},
    )
    assert r.json()["status"] == "modified" and r.json()["label"] == "correction"
    r = owner.post(
        f"/api/layers/{rel['id']}/relations",
        json={
            "from_span_id": spans["go"]["id"],
            "to_span_id": fine["id"],
            "label": "cause",
        },
    )
    assert r.status_code == 200 and r.json()["status"] == "confirmed"
    assert (
        owner.post(
            f"/api/layers/{rel['id']}/relations",
            json={
                "from_span_id": spans["go"]["id"],
                "to_span_id": spans["go"]["id"],
                "label": "cause",
            },
        ).status_code
        == 422
    )
    assert (
        owner.post(
            f"/api/layers/{layer['id']}/relations",
            json={
                "from_span_id": spans["go"]["id"],
                "to_span_id": fine["id"],
                "label": "cause",
            },
        ).status_code
        == 422
    )

    # agent: suggestions only, never on reviewed rows
    agent = bearer(owner, project)
    sent = items[4]
    r = agent.post(
        f"/api/layers/{layer['id']}/spans/suggest",
        json={
            "document_id": essay["document_id"],
            "sentence_index": 4,
            "spans": [
                {
                    "start_cp": sent["start_cp"],
                    "end_cp": sent["start_cp"] + 10,
                    "label": "句式",
                },
                {"start_cp": start, "end_cp": start + 4, "label": "时态"},
            ],
        },
    )
    assert r.status_code == 200, r.text
    out = r.json()
    assert [s["status"] for s in out["spans"]] == ["suggested", "modified"]
    assert out["spans"][0]["source"] == "model:agent" and out["skipped"] == [
        {"index": 1, "reason": "modified"}
    ]
    r = agent.post(
        f"/api/layers/{rel['id']}/spans/suggest",
        json={
            "document_id": essay["document_id"],
            "sentence_index": 4,
            "relations": [
                {
                    "from": str(out["spans"][0]["id"]),
                    "to": str(fine["id"]),
                    "label": "cause",
                }
            ],
        },
    )
    assert r.status_code == 200 and r.json()["relations"][0]["status"] == "suggested"
    # spans and relations in one call on the span layer: relations go to its only relation layer
    r = agent.post(
        f"/api/layers/{layer['id']}/spans/suggest",
        json={
            "document_id": essay["document_id"],
            "sentence_index": 4,
            "spans": [
                {
                    "start_cp": sent["start_cp"],
                    "end_cp": sent["start_cp"] + 3,
                    "label": "拼写",
                }
            ],
            "relations": [{"from": 0, "to": str(fine["id"]), "label": "correction"}],
        },
    )
    assert r.status_code == 200, r.text
    both = r.json()
    assert both["spans"][0]["label"] == "拼写" and both["skipped"] == []
    assert both["relations"][0]["from_span_id"] == both["spans"][0]["id"]
    assert len(sentences(owner, rel)["items"][4]["relations"]) == 2
    outside = agent.post(
        f"/api/layers/{layer['id']}/spans/suggest",
        json={
            "document_id": essay["document_id"],
            "sentence_index": 4,
            "spans": [{"start_cp": 0, "end_cp": 3, "label": "时态"}],
        },
    )
    assert outside.json()["error"]["code"] == "SPAN_OUT_OF_SENTENCE"
    assert (
        agent.post(
            f"/api/layers/{layer['id']}/spans",
            json={
                "document_id": essay["document_id"],
                "start_cp": 0,
                "end_cp": 3,
                "label": "时态",
            },
        ).status_code
        == 403
    )
    assert (
        agent.patch(
            f"/api/layers/{layer['id']}/spans/{out['spans'][0]['id']}",
            json={"action": "confirm"},
        ).status_code
        == 403
    )

    # sentence confirm: suggested spans of the sentence become confirmed and the sentence counts as reviewed
    r = owner.post(
        f"/api/layers/{layer['id']}/sentences/{essay['document_id']}/4/confirm", json={}
    )
    assert r.status_code == 200, r.text
    assert [s["status"] for s in r.json()["spans"]] == [
        "confirmed",
        "confirmed",
        "modified",
    ]


@pytest.fixture
def reviewed(owner, project, essay, fake_llm):
    """Span layer after review: go 时态 (confirmed), have 主谓一致 (confirmed), book 用词 (modified),
    was rejected, fine 时态 tense=past (human); relation have → book (correction)."""
    layer = span_layer(owner, project, essay)
    owner.post(
        f"/api/layers/{layer['id']}/preannotate",
        json={"sentence_range": {"start": 0, "end": 3}},
    )
    drain()
    spans = by_text(all_spans(owner, layer))
    patch_span(owner, layer, spans["go"], action="confirm")
    patch_span(owner, layer, spans["have"], action="confirm")
    patch_span(
        owner, layer, spans["book"], action="modify", attrs={"correction": "books"}
    )
    patch_span(owner, layer, spans["was"], action="reject")
    start = ESSAY.index("fine")
    owner.post(
        f"/api/layers/{layer['id']}/spans",
        json={
            "document_id": essay["document_id"],
            "start_cp": start,
            "end_cp": start + 4,
            "label": "时态",
            "attrs": {"tense": "past"},
        },
    )
    rel = make_layer(
        owner,
        project,
        kind="relation",
        name="更正",
        base_layer_id=layer["id"],
        label_set_id=label_set(owner, "error-correction")["id"],
    )
    relation = sentences(owner, rel)["items"][1]["relations"][0]
    owner.patch(
        f"/api/layers/{rel['id']}/relations/{relation['id']}",
        json={"action": "modify", "label": "correction"},
    )
    return layer, rel


def test_span_queries_match_sql(owner, project, essay, reviewed, fake_engine):
    layer, _rel = reviewed
    lid = layer["id"]
    ids = [x["id"] for x in label_set(owner, "learner-errors")["labels"]]
    cases = [
        ("[时态]", ["时态"], None),
        ("[偏误:时态]", ["时态"], None),
        ("[偏误]", ids, None),
        ("[主谓一致]", ["主谓一致"], None),
        ("[时态:tense=past]", ["时态"], {"tense": "past"}),
        ("[用词:correction=books]", ["用词"], {"correction": "books"}),
        ("[偏误:用词]", ["用词"], None),
    ]
    for q, labels, attrs in cases:
        res = search(
            owner, essay, q, layers={"span": lid}, context={"unit": "sentence"}
        )
        assert res["counts"]["hits"] == effective_count(lid, labels, attrs), q
    res = search(
        owner, essay, "[时态]", layers={"span": lid}, context={"unit": "sentence"}
    )
    assert [h["match"] for h in res["items"]] == ["go", "fine"]
    assert res["items"][0]["left"] == "Yesterday I " and res["counts"]["sentences"] == 2
    assert res["parse"]["tokens"][0]["kind"] == "span"
    # default layer: the corpus's span layer
    assert search(owner, essay, "[时态]")["counts"]["hits"] == 2
    bad = owner.post(
        "/api/search",
        json={
            "scope": {"corpus_id": essay["id"]},
            "query": "[不存在]",
            "layers": {"span": lid},
        },
    )
    assert bad.json()["error"]["code"] == "INVALID_LABEL"
    # anchored: the word's token must lie inside a matching span
    fake_engine.needle = "go"
    assert search(owner, essay, "go[时态]", layers={"span": lid})["counts"]["hits"] == 1
    assert (
        search(owner, essay, "go[主谓一致]", layers={"span": lid})["counts"]["hits"]
        == 0
    )
    fake_engine.needle = "was"
    assert (
        search(owner, essay, "was[主谓一致]", layers={"span": lid})["counts"]["hits"]
        == 0
    )

    # statistics: label frequency over effective spans and relations
    r = owner.post(
        "/api/stats/frequency",
        json={
            "scope": {"corpus_id": essay["id"]},
            "unit": "span_label",
            "layers": {"span": lid},
        },
    )
    assert r.status_code == 200, r.text
    drain()
    stat_id = owner.get(f"/api/jobs/{r.json()['id']}").json()["result"]["stat_id"]
    rows = owner.get(f"/api/stats/{stat_id}").json()["rows"]
    expected = dict(
        sql(
            "SELECT label, count(*) FROM span_annotations WHERE layer_id = CAST(:l AS uuid)"
            " AND status IN ('confirmed', 'modified') GROUP BY label",
            l=lid,
        )
    )
    assert {r["key"]: r["values"]["freq"] for r in rows} == expected
    assert rows[0]["search_query"] == "[时态]"
    r = owner.post(
        "/api/stats/frequency",
        json={"scope": {"corpus_id": essay["id"]}, "unit": "relation_label"},
    )
    drain()
    stat_id = owner.get(f"/api/jobs/{r.json()['id']}").json()["result"]["stat_id"]
    rows = owner.get(f"/api/stats/{stat_id}").json()["rows"]
    assert {r["key"]: r["values"]["freq"] for r in rows} == {"correction": 1}
    bad = owner.post(
        "/api/stats/collocations",
        json={"scope": {"corpus_id": essay["id"]}, "node": "go[时态]"},
    )
    assert bad.status_code == 422


def export(client, layer, fmt):
    r = client.post(f"/api/layers/{layer['id']}/export", json={"format": fmt})
    assert r.status_code == 200, r.text
    drain()
    job = client.get(f"/api/jobs/{r.json()['id']}").json()
    assert job["state"] == "succeeded", job
    d = client.get(f"/api/jobs/{job['id']}/download")
    assert d.status_code == 200
    return job, d


def test_span_exports(owner, project, essay, reviewed):
    layer, rel = reviewed
    job, d = export(owner, layer, "brat")
    assert d.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(d.content))
    names = sorted(z.namelist())
    assert names == ["001-essay.ann", "001-essay.txt", "annotation.conf"]
    txt = z.read("001-essay.txt").decode("utf-8")
    assert txt == ESSAY
    lines = z.read("001-essay.ann").decode("utf-8").splitlines()
    spans = effective_count(
        layer["id"], [x["id"] for x in label_set(owner, "learner-errors")["labels"]]
    )
    assert (
        len(lines) == spans + 1 == job["result"]["spans"] + job["result"]["relations"]
    )
    t_lines = [x.split("\t") for x in lines if x.startswith("T")]
    for _, kind_range, body in t_lines:
        _label, a, b = kind_range.split(" ")
        assert txt[int(a) : int(b)] == body
    labels = {body: kind_range.split(" ")[0] for _, kind_range, body in t_lines}
    assert labels == {"go": "时态", "have": "主谓一致", "book": "用词", "fine": "时态"}
    r_line = next(x for x in lines if x.startswith("R"))
    ids = {body: tid for tid, _, body in t_lines}
    assert r_line == f"R1\tcorrection Arg1:{ids['have']} Arg2:{ids['book']}"
    conf = z.read("annotation.conf").decode("utf-8")
    assert "\ncorrection\tArg1:<ENTITY>, Arg2:<ENTITY>\n" in conf

    job, d = export(owner, layer, "xlsx")
    wb = load_workbook(io.BytesIO(d.content))
    ws = wb["偏误分布"]
    header = [c.value for c in ws[1]]
    assert header == ["标签", "essay", "合计"]
    assert all(c.font.name == "宋体" and c.font.b for c in ws[1])
    assert all(c.fill.fill_type is None for row in ws.iter_rows() for c in row)
    counts = {row[0].value: row[1].value for row in ws.iter_rows(min_row=2)}
    assert counts["时态"] == 2 and counts["主谓一致"] == 1 and counts["用词"] == 1
    assert counts["拼写"] == 0 and counts["合计"] == 4
    span_rows = list(wb["片段"].iter_rows(min_row=2, values_only=True))
    assert sorted(r[4] for r in span_rows) == ["book", "fine", "go", "have"]

    job, d = export(owner, rel, "json")
    data = json.loads(d.content)
    assert data["documents"][0]["text"] == ESSAY and data["offsets"] == "code points"
    kinds = {x["kind"]: x for x in data["layers"]}
    assert len(kinds["span"]["spans"]) == 4 and len(kinds["relation"]["relations"]) == 1
    assert (
        owner.post(
            f"/api/layers/{layer['id']}/export", json={"format": "tmx"}
        ).status_code
        == 422
    )


# ---- multiple annotators ----


def kappa(a, b):
    """Independent Cohen's κ."""
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] / n * cb[k] / n for k in set(a) | set(b))
    return (po - pe) / (1 - pe)


def alpha_two(a, b):
    """Independent nominal α for two coders without missing values (Krippendorff 2011, section A)."""
    values = a + b
    n = len(values)
    nc = Counter(values)
    disagree = sum(x != y for x, y in zip(a, b)) * 2  # ordered pairs
    expected = sum(nc[c] * nc[k] for c in nc for k in nc if c != k)
    return 1 - (n - 1) * disagree / expected


def test_multi_annotator_spans(owner, project, essay, fake_llm, member, make_user):
    ann = member("annotator")
    layer = span_layer(
        owner, project, essay, mode="multi", annotators=[owner.user_id, ann.user_id]
    )
    lid = layer["id"]
    owner.post(
        f"/api/layers/{lid}/preannotate",
        json={"sentence_range": {"start": 0, "end": 3}},
    )
    drain()
    layer = owner.get(f"/api/layers/{lid}").json()
    assert (
        layer["mode"] == "multi" and layer["progress"]["total"] == 8
    )  # four suggestions each
    mine = all_spans(ann, layer)
    assert {s["annotator_id"] for s in mine} == {ann.user_id} and len(mine) == 4
    assert {s["annotator_id"] for s in all_spans(owner, layer)} == {owner.user_id}
    theirs = all_spans(owner, layer, annotator_id=ann.user_id)
    assert {s["annotator_id"] for s in theirs} == {ann.user_id}
    assert (
        ann.get(
            f"/api/layers/{lid}/spans", params={"annotator_id": owner.user_id}
        ).status_code
        == 403
    )
    outsider = member("editor")
    assert len(all_spans(outsider, layer)) == 8  # editor not listed: sees everyone
    assert (
        outsider.patch(
            f"/api/layers/{lid}/spans/{mine[0]['id']}", json={"action": "confirm"}
        ).status_code
        == 403
    )
    # an annotator cannot act on another annotator's rows
    other = by_text(all_spans(owner, layer))["go"]
    assert (
        ann.patch(
            f"/api/layers/{lid}/spans/{other['id']}", json={"action": "confirm"}
        ).status_code
        == 404
    )

    doc = essay["document_id"]
    for s in range(3):
        assert (
            owner.post(
                f"/api/layers/{lid}/sentences/{doc}/{s}/confirm", json={}
            ).status_code
            == 200
        )
    a = by_text(mine)
    patch_span(ann, layer, a["have"], action="modify", label="时态")
    patch_span(ann, layer, a["was"], action="reject")
    for s in range(3):
        ann.post(f"/api/layers/{lid}/sentences/{doc}/{s}/confirm", json={})

    # before adjudication only agreed spans are effective
    assert (
        search(owner, essay, "[时态]", layers={"span": lid})["counts"]["hits"] == 1
    )  # go
    assert (
        search(owner, essay, "[主谓一致]", layers={"span": lid})["counts"]["hits"] == 0
    )

    agreement = owner.get(f"/api/layers/{lid}/agreement").json()
    assert {x["user_id"]: x["done"] for x in agreement["annotators"]} == {
        owner.user_id: 3,
        ann.user_id: 3,
    }
    # units (extent order): go, have, book, was
    o = ["时态", "主谓一致", "用词", "主谓一致"]
    b = ["时态", "时态", "用词", "∅"]
    pair = agreement["pairwise"][0]
    assert pair["n"] == 4 and pair["kappa"] == pytest.approx(kappa(o, b))
    assert agreement["alpha"] == pytest.approx(alpha_two(o, b))
    assert agreement["lenient"]["n"] == 4
    assert {x["label"] for x in agreement["per_label"]} == {
        "时态",
        "主谓一致",
        "用词",
        "∅",
    }

    dis = owner.get(f"/api/layers/{lid}/disagreements").json()
    assert dis["total"] == 2 and not any(x["adjudicated"] for x in dis["items"])
    by_word = {
        ESSAY[int(x["target_key"].rsplit(":", 1)[1].split("-")[0]) :].split(" ")[0]: x
        for x in dis["items"]
    }
    assert set(by_word) == {"have", "was"}
    values = {v["annotator_id"]: v["value"] for v in by_word["have"]["values"]}
    assert (
        values[owner.user_id]["label"] == "主谓一致"
        and values[ann.user_id]["label"] == "时态"
    )
    assert (
        by_word["was"]["sentence_index"] == 2
        and by_word["was"]["text"] == "We was happy ."
    )
    assert ann.get(f"/api/layers/{lid}/disagreements").status_code == 403

    r = owner.post(
        f"/api/layers/{lid}/adjudicate",
        json={
            "target_kind": "span",
            "target_key": by_word["have"]["target_key"],
            "final": {"label": "时态"},
        },
    )
    assert r.status_code == 200, r.text
    assert r.json()["final"] == {"label": "时态", "attrs": {}}
    owner.post(
        f"/api/layers/{lid}/adjudicate",
        json={
            "target_kind": "span",
            "target_key": by_word["was"]["target_key"],
            "final": {"label": None},
        },
    )
    assert (
        ann.post(
            f"/api/layers/{lid}/adjudicate",
            json={
                "target_kind": "span",
                "target_key": by_word["was"]["target_key"],
                "final": {"label": None},
            },
        ).status_code
        == 403
    )
    assert (
        owner.post(
            f"/api/layers/{lid}/adjudicate",
            json={
                "target_kind": "token",
                "target_key": "x:1:1",
                "final": {"label": None},
            },
        ).status_code
        == 422
    )
    res = search(owner, essay, "[时态]", layers={"span": lid})
    assert res["counts"]["hits"] == 2 and [h["match"] for h in res["items"]] == [
        "go",
        "have",
    ]
    assert all(
        x["adjudicated"]
        for x in owner.get(f"/api/layers/{lid}/disagreements").json()["items"]
    )
    _, d = export(owner, layer, "brat")
    ann_lines = (
        zipfile.ZipFile(io.BytesIO(d.content))
        .read("001-essay.ann")
        .decode()
        .splitlines()
    )
    assert len(ann_lines) == 3  # go, have (adjudicated), book

    # not listed in a multi layer → cannot annotate
    other_ann = member("viewer")
    assert (
        other_ann.post(
            f"/api/layers/{lid}/spans",
            json={"document_id": doc, "start_cp": 0, "end_cp": 3, "label": "时态"},
        ).status_code
        == 403
    )


def test_multi_annotator_tokens(owner, project, essay, member):
    ann = member("annotator")
    pos = make_layer(
        owner,
        project,
        corpus_id=essay["id"],
        kind="pos",
        name="词性",
        label_set_id=label_set(owner, "upos")["id"],
        mode="multi",
        annotators=[owner.user_id, ann.user_id],
    )
    lid = pos["id"]
    doc = essay["document_id"]
    tokens = len(ESSAY.split())
    assert pos["progress"]["total"] == 2 * tokens
    s0 = ann.get(f"/api/layers/{lid}/sentences/{doc}/0").json()
    assert {a["annotator_id"] for a in s0["annotations"]} == {ann.user_id}

    def act(client, t, label):
        r = client.post(
            f"/api/layers/{lid}/annotations",
            json={
                "document_id": doc,
                "sentence_index": 0,
                "token_index": t,
                "action": "modify",
                "label": label,
            },
        )
        assert r.status_code == 200, r.text

    act(owner, 2, "PRON")
    act(ann, 2, "PRON")
    act(owner, 3, "VERB")
    act(ann, 3, "AUX")
    for client in (owner, ann):
        client.post(f"/api/layers/{lid}/sentences/{doc}/0/confirm", json={})
    agreement = owner.get(f"/api/layers/{lid}/agreement").json()
    n = len(ESSAY.split("\n")[0].split())
    assert agreement["pairwise"][0]["n"] == n
    o = [None] * n
    b = [None] * n
    o[1], b[1], o[2], b[2] = "PRON", "PRON", "VERB", "AUX"
    assert agreement["pairwise"][0]["kappa"] == pytest.approx(kappa(o, b))

    def pos_hits(tag):
        r = owner.post(
            "/api/search",
            json={
                "scope": {"corpus_id": essay["id"]},
                "query": f"{{pos={tag}}}",
                "layers": {"pos": lid},
            },
        )
        assert r.status_code == 200, r.text
        return r.json()["counts"]["hits"]

    assert pos_hits("PRON") == 1 and pos_hits("VERB") == 0 and pos_hits("AUX") == 0
    dis = owner.get(f"/api/layers/{lid}/disagreements").json()["items"]
    assert [d["target_key"] for d in dis] == [f"{doc}:0:3"]
    owner.post(
        f"/api/layers/{lid}/adjudicate",
        json={
            "target_kind": "token",
            "target_key": f"{doc}:0:3",
            "final": {"label": "VERB"},
        },
    )
    assert pos_hits("VERB") == 1


# ---- alignment ----


@pytest.fixture
def tmx(owner, project):
    c = owner.post(
        f"/api/projects/{project['id']}/corpora",
        json={"name": "平行", "language": "en"},
    ).json()
    upload(owner, c["id"], {"t.tmx": "\n".join(f"{a}\t{b}" for a, b in TMX)})
    drain()
    docs = owner.get(f"/api/corpora/{c['id']}/documents").json()["items"]
    c["en"] = next(d["id"] for d in docs if d["language"] == "en")
    c["zh"] = next(d["id"] for d in docs if d["language"] == "zh")
    return c


def pairs(client, layer, **params):
    r = client.get(f"/api/layers/{layer['id']}/alignments", params=params)
    assert r.status_code == 200, r.text
    return r.json()


def test_alignment_layers(owner, project, tmx, fake_llm):
    sent = make_layer(
        owner,
        project,
        corpus_id=tmx["id"],
        kind="alignment",
        name="句对齐",
        level="sentence",
    )
    assert sent["status"] == "ready" and sent["tgt_corpus_id"] == tmx["id"]
    assert sent["src_corpus_id"] == tmx["id"] and sent["documents"]["total"] == 1
    assert sent["source_model"] == "sentence_alignments" and fake_llm["calls"] == []
    items = pairs(owner, sent)["items"]
    assert len(items) == 3 and all(
        p["alignments"][0]["source"] == "system" for p in items
    )
    assert (
        items[1]["src"]["text"] == "The cat sleeps ."
        and items[1]["tgt"]["text"] == "猫 在 睡觉 。"
    )
    assert [t["form"] for t in items[0]["tgt"]["tokens"]] == [
        "我",
        "喜欢",
        "苹果",
        "。",
    ]
    page = pairs(owner, sent, page_size=2)
    assert len(page["items"]) == 2 and page["next_cursor"] == f"{tmx['en']}:1"
    assert len(pairs(owner, sent, cursor=page["next_cursor"])["items"]) == 1
    stats = owner.post("/api/stats/alignment", json={"layer_id": sent["id"]}).json()
    assert (
        stats["pairs"] == 3
        and stats["one_to_one"] == 3
        and stats["unaligned"] == {"src": 0, "tgt": 0}
    )
    # move a target sentence (imported pairs stay effective until changed)
    first = items[2]["alignments"][0]
    r = owner.patch(
        f"/api/layers/{sent['id']}/alignments/{first['id']}",
        json={"action": "modify", "tgt_sentence_index": 1},
    )
    assert r.status_code == 200 and r.json()["status"] == "modified"
    stats = owner.post("/api/stats/alignment", json={"layer_id": sent["id"]}).json()
    assert stats["one_to_one"] == 1 and stats["one_to_many"] == 1
    assert stats["unaligned"] == {"src": 0, "tgt": 1}
    _, d = export(owner, sent, "tmx")
    tus = etree.fromstring(d.content).findall(".//tu")
    assert len(tus) == 3
    segs = [[s.text for s in tu.findall(".//seg")] for tu in tus]
    assert segs[2] == ["Good night .", "猫 在 睡觉 。"]

    word = make_layer(
        owner,
        project,
        corpus_id=tmx["id"],
        kind="alignment",
        name="词对齐",
        level="word",
    )
    assert word["source_model"] == "sentence_alignments + fake-chat"
    assert len(fake_llm["calls"]) == 1  # three pairs in one request
    items = pairs(owner, word)["items"]
    links = [a for a in items[0]["alignments"] if a["level"] == "word"]
    assert [(a["src_tokens"], a["tgt_tokens"]) for a in links] == [
        ([1], [1]),
        ([2], [2]),
        ([3], [3]),
        ([4], [4]),
    ]
    assert {a["status"] for a in links} == {"suggested"}
    job = owner.get(f"/api/jobs/{word['job_id']}").json()
    assert job["result"]["word_dropped"] == 1

    # review: confirm pair 0; in pair 1 relink cat (2) → 猫 (1), drop the rest by confirming the sentence
    r = owner.post(f"/api/layers/{word['id']}/sentences/{tmx['en']}/0/confirm", json={})
    assert r.status_code == 200 and {
        a["status"] for a in r.json()["items"][0]["alignments"]
    } == {"confirmed"}
    second = [a for a in items[1]["alignments"] if a["level"] == "word"]
    owner.patch(
        f"/api/layers/{word['id']}/alignments/{second[0]['id']}",
        json={"action": "modify", "src_tokens": [2], "tgt_tokens": [1]},
    )
    owner.patch(
        f"/api/layers/{word['id']}/alignments/{second[1]['id']}",
        json={"action": "reject"},
    )
    r = owner.post(
        f"/api/layers/{word['id']}/alignments",
        json={
            "src_document_id": tmx["en"],
            "src_sentence_index": 1,
            "tgt_document_id": tmx["zh"],
            "tgt_sentence_index": 1,
            "src_tokens": [3],
            "tgt_tokens": [2, 3],
        },
    )
    assert r.status_code == 200 and r.json()["level"] == "word"
    bad = owner.post(
        f"/api/layers/{word['id']}/alignments",
        json={
            "src_document_id": tmx["en"],
            "src_sentence_index": 1,
            "tgt_document_id": tmx["zh"],
            "tgt_sentence_index": 1,
            "src_tokens": [30],
            "tgt_tokens": [1],
        },
    )
    assert bad.json()["error"]["code"] == "INVALID_TOKENS"
    for a in [
        a
        for a in pairs(owner, word)["items"][1]["alignments"]
        if a["status"] == "suggested"
    ]:
        owner.patch(
            f"/api/layers/{word['id']}/alignments/{a['id']}", json={"action": "confirm"}
        )
    # agent suggestion on pair 2
    agent = bearer(owner, project)
    r = agent.post(
        f"/api/layers/{word['id']}/alignments/suggest",
        json={
            "src_document_id": tmx["en"],
            "src_sentence_index": 2,
            "tgt_document_id": tmx["zh"],
            "tgt_sentence_index": 2,
            "links": [{"src": [1, 2], "tgt": [1]}, {"src": [1], "tgt": [1]}],
        },
    )
    assert r.status_code == 200, r.text
    assert [a["src_tokens"] for a in r.json()["items"]] == [[1, 2]]
    assert r.json()["skipped"][0]["reason"] == "suggested"
    assert (
        agent.post(
            f"/api/layers/{word['id']}/alignments/suggest",
            json={
                "src_document_id": tmx["en"],
                "src_sentence_index": 2,
                "tgt_document_id": tmx["zh"],
                "tgt_sentence_index": 2,
                "links": [{"src": [9], "tgt": [1]}],
            },
        ).status_code
        == 422
    )

    _, d = export(owner, word, "tmx")
    root = etree.fromstring(d.content)
    props = [
        tu.find("prop[@type='x-word-alignment']").text for tu in root.findall(".//tu")
    ]
    assert props[0] == "1-1 2-2 3-3 4-4"
    assert props[1] == "2-1 3-2 3-3 4-4"  # modified, rejected, added
    assert props[2] is None  # only suggestions: not effective
    stats = owner.post("/api/stats/alignment", json={"layer_id": word["id"]}).json()
    assert stats["pairs"] == 3 and stats["avg_word_pairs"] == pytest.approx(
        (4 + 4 + 0) / 3
    )


def test_alignment_candidates_without_stored_pairs(owner, project, fake_llm):
    pid = project["id"]
    en = owner.post(
        f"/api/projects/{pid}/corpora", json={"name": "EN", "language": "en"}
    ).json()
    zh = owner.post(
        f"/api/projects/{pid}/corpora", json={"name": "ZH", "language": "zh"}
    ).json()
    upload(owner, en["id"], {"story.txt": "\n".join(a for a, _ in TMX)})
    upload(owner, zh["id"], {"story.txt": "\n".join([TMX[1][1], TMX[0][1], TMX[2][1]])})
    drain()
    layer = make_layer(
        owner,
        project,
        corpus_id=en["id"],
        tgt_corpus_id=zh["id"],
        kind="alignment",
        name="候选",
        level="sentence",
    )
    assert layer["status"] == "ready" and layer["source_model"] == "MiniLM"
    items = pairs(owner, layer)["items"]
    got = [(p["src"]["sentence_index"], p["tgt"]["sentence_index"]) for p in items]
    assert got == [(0, 1), (1, 0), (2, 2)]
    assert all(
        a["source"] == "model:minilm" and a["score"] >= 0.5
        for p in items
        for a in p["alignments"]
    )
    stats = owner.post("/api/stats/alignment", json={"layer_id": layer["id"]}).json()
    assert stats["pairs"] == 0  # candidates are suggestions until confirmed


# ---- review progress, single-mode writers, export attribution ----


def test_layer_documents_done(owner, project, essay, fake_llm, member):
    """A sentence is done once confirmed or when its rows are all reviewed; a document is done (done = sentences)
    when it has rows and none is suggested. Single-mode layers are written only by their creator."""
    layer = span_layer(owner, project, essay)
    lid, doc = layer["id"], essay["document_id"]

    def progress(client=owner):
        item = client.get(f"/api/layers/{lid}/documents").json()["items"][0]
        return item["sentences"], item["done"], item["pending"]

    owner.post(
        f"/api/layers/{lid}/preannotate",
        json={"sentence_range": {"start": 0, "end": 3}},
    )
    drain()
    assert progress() == (5, 0, 3)  # suggestions in sentences 0–2 only
    r = owner.post(f"/api/layers/{lid}/sentences/{doc}/4/confirm", json={})
    assert r.status_code == 200, r.text
    assert progress() == (5, 1, 3)
    spans = by_text(all_spans(owner, layer))
    patch_span(owner, layer, spans["go"], action="confirm")
    assert progress() == (5, 2, 2)

    annotator, editor = member("annotator"), member("editor")
    for client in (annotator, editor):
        r = client.patch(
            f"/api/layers/{lid}/spans/{spans['have']['id']}", json={"action": "confirm"}
        )
        assert r.status_code == 403 and r.json()["error"]["code"] == "NOT_ANNOTATOR"
        assert progress(client) == (5, 2, 2)

    patch_span(owner, layer, spans["have"], action="confirm")
    patch_span(owner, layer, spans["book"], action="reject")
    patch_span(owner, layer, spans["was"], action="modify", label="时态")
    assert progress() == (5, 5, 0)  # sentence 3 never reviewed, but no row is pending
    assert owner.get(f"/api/layers/{lid}").json()["documents"] == {
        "total": 1,
        "done": 1,
    }


def test_layer_exports_carry_source_notice(owner, project, essay, reviewed, tmx):
    meta = {"license": "CC BY 4.0", "attribution": "某某大学"}
    sql(
        "UPDATE corpora SET metadata = CAST(:m AS jsonb) WHERE id IN"
        " (CAST(:a AS uuid), CAST(:b AS uuid)) RETURNING id",
        m=json.dumps(meta, ensure_ascii=False),
        a=essay["id"],
        b=tmx["id"],
    )
    notice = "来源：作文；许可：CC BY 4.0；署名：某某大学"
    layer, rel = reviewed

    _, d = export(owner, layer, "brat")
    z = zipfile.ZipFile(io.BytesIO(d.content))
    assert "README.txt" in z.namelist()
    assert z.read("README.txt").decode("utf-8") == notice + "\n"
    conf = z.read("annotation.conf").decode("utf-8")
    assert conf.startswith(f"# {notice}\n[entities]\n")

    _, d = export(owner, layer, "xlsx")
    wb = load_workbook(io.BytesIO(d.content))
    for name, first in (("偏误分布", "标签"), ("片段", "文档")):
        rows = list(wb[name].iter_rows(values_only=True))
        assert rows[0][0] == notice and not any(rows[1]) and rows[2][0] == first

    _, d = export(owner, rel, "json")
    data = json.loads(d.content)
    assert data["source"] == [notice] and next(iter(data)) == "source"

    sent = make_layer(
        owner,
        project,
        corpus_id=tmx["id"],
        kind="alignment",
        name="句对齐",
        level="sentence",
    )
    _, d = export(owner, sent, "tmx")
    root = etree.fromstring(d.content)
    assert [n.text for n in root.find("header").findall("note")] == [
        "来源：平行；许可：CC BY 4.0；署名：某某大学"
    ]
    assert len(root.findall(".//tu")) == 3
