"""P2 api: stats jobs (engine.stats stubbed), evidence, agent tokens, temporary-filter scope."""

import csv
import io
import json
from datetime import UTC, datetime, timedelta

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from tests.api.conftest import drain, upload

TEXT_A = "Language is a system.\nThe language of science."
TEXT_B = "Grammar rules apply.\nA language again."
COLUMNS = [
    {"key": "item", "label": "项", "type": "string"},
    {"key": "freq", "label": "频次", "type": "number"},
]
TOTALS = {"tokens": 14, "sentences": 4, "documents": 2}


@pytest.fixture
def corpus(owner, project):
    c = owner.post(
        f"/api/projects/{project['id']}/corpora",
        json={"name": "作文", "language": "en"},
    ).json()
    meta = [{"year": 2001, "genre": "news"}, {"year": 2010, "genre": "fiction"}]
    upload(owner, c["id"], {"a.txt": TEXT_A}, metadata=json.dumps(meta[0]))
    upload(owner, c["id"], {"b.txt": TEXT_B}, metadata=json.dumps(meta[1]))
    drain()
    return c


@pytest.fixture
def fake_stats(monkeypatch):
    from engine import stats

    calls = []

    def fake(method):
        def run(*args, **kw):
            calls.append({"method": method, "args": args, "kw": kw})
            if method == "distribution":
                rows = [
                    {
                        "key": str(y),
                        "values": {"bin": str(y), "hits": 2},
                        "search_query": args[1],
                        "bin_value": y,
                    }
                    for y in (2001, 2010)
                ]
            else:
                rows = [
                    {
                        "key": w,
                        "values": {"item": w, "freq": 9 - i},
                        "search_query": f'"{w}"',
                    }
                    for i, w in enumerate(["language", "is", "a"])
                ]
            return stats.Table(COLUMNS, rows, 7, TOTALS)

        return run

    for m in ("frequency", "ngrams", "collocations", "keywords", "distribution"):
        monkeypatch.setattr(stats, m, fake(m))
    return calls


def run_stat(client, method, body):
    r = client.post(f"/api/stats/{method}", json=body)
    assert r.status_code == 200, r.text
    job = r.json()
    assert job["kind"] == "stats" and job["state"] == "queued"
    drain()
    job = client.get(f"/api/jobs/{job['id']}").json()
    assert job["state"] == "succeeded", job
    return job["result"]["stat_id"]


def test_stats_job_lifecycle(owner, project, corpus, member, fake_stats):
    scope = {"corpus_id": corpus["id"]}
    stat_id = run_stat(owner, "frequency", {"scope": scope, "unit": "word", "top": 2})
    assert fake_stats[-1]["args"][1:3] == ("word", None)
    assert fake_stats[-1]["kw"] == {"exclude_punct": True}  # punctuation left out by default
    stat = owner.get(f"/api/stats/{stat_id}").json()
    assert stat["method"] == "frequency" and stat["project_id"] == project["id"]
    assert stat["scope"] == {
        "corpus_id": corpus["id"],
        "subcorpus_id": None,
        "name": "作文",
    }
    assert stat["params"]["top"] == 2 and stat["params"]["unit"] == "word"
    assert stat["totals"] == TOTALS and stat["columns"] == COLUMNS
    assert stat["row_count"] == 7 and stat["index_version"]
    assert [r["key"] for r in stat["rows"]] == ["language", "is"]  # capped at top
    assert stat["rows"][0]["search_scope"] == scope
    assert stat["rows"][0]["search_query"] == '"language"'

    listed = owner.get(f"/api/projects/{project['id']}/stats").json()
    assert listed["total"] == 1 and "rows" not in listed["items"][0]

    viewer = member("viewer")
    assert viewer.get(f"/api/stats/{stat_id}").status_code == 200
    job = viewer.post(f"/api/stats/{stat_id}/export", json={"format": "csv"}).json()
    assert job["kind"] == "export" and job["project_id"] == project["id"]
    drain()
    r = viewer.get(f"/api/jobs/{job['id']}/download")
    assert r.status_code == 200, r.text
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    # parameters, scope and token total (punctuation included) above the table
    assert rows[0] == ["统计：频次（单位 词形；最低频次 1；显示前 2；排除标点和数字）"]
    assert rows[1] == ["范围：作文，14 词（含标点）"]
    assert rows[2][0].startswith("索引版本：")
    head = rows.index(["项", "频次", "查询式"])
    assert rows[head + 1 :] == [
        ["language", "9", '"language"'],
        ["is", "8", '"is"'],
        ["a", "7", '"a"'],
    ]
    assert job["summary"].startswith("频次")
    assert viewer.get(f"/api/jobs/{job['id']}").json()["result"]["filename"].startswith(
        "统计-频次-"
    )
    assert viewer.delete(f"/api/stats/{stat_id}").status_code == 403

    assert owner.delete(f"/api/stats/{stat_id}").status_code == 204
    assert owner.get(f"/api/stats/{stat_id}").status_code == 404


def test_stats_other_methods_and_validation(owner, project, corpus, fake_stats):
    scope = {"corpus_id": corpus["id"]}
    run_stat(owner, "ngrams", {"scope": scope, "n": 3, "pos_pattern": ["v", None, "n"]})
    run_stat(
        owner,
        "ngrams",
        {"scope": scope, "n": 2, "contains": "language", "exclude_punct": False},
    )
    assert fake_stats[-1]["kw"] == {"contains": "language", "exclude_punct": False}
    run_stat(owner, "collocations", {"scope": scope, "node": "language", "pos": "n"})
    assert fake_stats[-1]["args"][1:3] == ("language", 5)
    sub = owner.post(
        f"/api/projects/{project['id']}/subcorpora",
        json={"name": "新闻", "corpus_id": corpus["id"], "filter": {"genre": "news"}},
    ).json()
    stat_id = run_stat(
        owner, "keywords", {"target": {"subcorpus_id": sub["id"]}, "reference": scope}
    )
    target, reference = fake_stats[-1]["args"][:2]
    assert "genre" in target.restrict and reference.restrict is None
    stat = owner.get(f"/api/stats/{stat_id}").json()
    assert stat["scope"]["name"] == "新闻"
    assert stat["params"]["reference"]["name"] == "作文"
    assert stat["rows"][0]["search_scope"] == {"subcorpus_id": sub["id"]}

    bad = [
        ("frequency", {"scope": scope, "pos": "zz"}, "VALIDATION_ERROR"),
        ("collocations", {"scope": scope, "node": "a $"}, "QUERY_INVALID"),
        ("distribution", {"scope": scope, "field": "nope"}, "INVALID_FIELD"),
        ("ngrams", {"scope": scope, "n": 2, "pos_pattern": ["v"]}, "VALIDATION_ERROR"),
        ("keywords", {"target": scope, "reference": scope}, "VALIDATION_ERROR"),
        (
            "collocations",
            {"scope": scope, "node": "x", "measures": ["dice"]},
            "VALIDATION_ERROR",
        ),
    ]
    for method, body, code in bad:
        r = owner.post(f"/api/stats/{method}", json=body)
        assert r.status_code == 422 and r.json()["error"]["code"] == code, (
            method,
            r.text,
        )


def test_distribution_rows_use_temporary_filter(owner, project, corpus, fake_stats):
    sub = owner.post(
        f"/api/projects/{project['id']}/subcorpora",
        json={
            "name": "新闻",
            "corpus_id": corpus["id"],
            "filter": {"genre": ["news", "fiction"]},
        },
    ).json()
    stat_id = run_stat(
        owner,
        "distribution",
        {"scope": {"subcorpus_id": sub["id"]}, "query": "language", "field": "year"},
    )
    call = fake_stats[-1]
    assert call["args"][1:3] == ("language", "year") and call["kw"]["numeric"] is True
    rows = owner.get(f"/api/stats/{stat_id}").json()["rows"]
    assert "bin_value" not in rows[0]
    assert rows[0]["search_scope"] == {
        "corpus_id": corpus["id"],
        "filter": {"genre": ["news", "fiction"], "year": 2001},
    }
    # the row's scope is accepted by search as a temporary filter
    r = owner.post(
        "/api/search",
        json={"scope": rows[0]["search_scope"], "query": rows[0]["search_query"]},
    )
    assert r.status_code == 200, r.text


def test_search_temporary_filter_scope(owner, project, corpus, fake_engine):
    r = owner.post(
        "/api/search",
        json={
            "scope": {"corpus_id": corpus["id"], "filter": {"year": {"gte": 2005}}},
            "query": "language",
        },
    )
    assert r.status_code == 200, r.text
    assert "int(" in fake_engine.calls[-1]["restrict"]
    r = owner.post(
        "/api/search",
        json={
            "scope": {"corpus_id": corpus["id"], "filter": {"nope": 1}},
            "query": "x",
        },
    )
    assert r.status_code == 422 and r.json()["error"]["code"] == "INVALID_FILTER"
    sub = owner.post(
        f"/api/projects/{project['id']}/subcorpora",
        json={"name": "s", "corpus_id": corpus["id"], "filter": {"year": 2001}},
    ).json()
    r = owner.post(
        "/api/search",
        json={
            "scope": {"subcorpus_id": sub["id"], "filter": {"year": 2001}},
            "query": "x",
        },
    )
    assert r.status_code == 422


def test_evidence_crud_and_export(owner, project, corpus, member, make_user):
    hit = owner.post(
        "/api/search", json={"scope": {"corpus_id": corpus["id"]}, "query": "language"}
    ).json()
    h = hit["items"][0]
    body = {
        "kind": "hit",
        "ref": {
            "search_query": "language",
            "scope": {"corpus_id": corpus["id"]},
            "hit_id": h["hit_id"],
            "document_id": h["document_id"],
            "start_cp": h["start_cp"],
            "end_cp": h["end_cp"],
        },
        "text": "Language is a system.",
        "source": "作文 / a",
        "note": "定义句",
    }
    r = owner.post(f"/api/projects/{project['id']}/evidence", json=body)
    assert r.status_code == 200, r.text
    ev = r.json()
    assert ev["ref"]["hit_id"] == h["hit_id"] and ev["created_by"] == owner.user_id
    # a pinned hit is stored as its own sentence with the match marked, not the client's KWIC text
    assert ev["text"] == "**Language** is a system."
    assert ev["source"] == "a · 句 1 · 作文"
    assert ev["ref"]["sentence"] == "Language is a system."
    assert ev["ref"]["match"] == [0, 8] and ev["ref"]["sentence_index"] == 0
    assert ev["ref"]["aligned"] == [] and ev["ref"]["document_title"] == "a"
    annotator = member("annotator")
    own = annotator.post(
        f"/api/projects/{project['id']}/evidence",
        json={"kind": "text", "text": "随手记", "source": "笔记"},
    ).json()
    viewer = member("viewer")
    assert (
        viewer.post(f"/api/projects/{project['id']}/evidence", json=body).status_code
        == 403
    )
    bad = {**body, "ref": {"stat_id": "00000000-0000-0000-0000-000000000000"}}
    assert (
        owner.post(f"/api/projects/{project['id']}/evidence", json=bad).status_code
        == 404
    )
    assert (
        make_user("x@example.edu")
        .get(f"/api/projects/{project['id']}/evidence")
        .status_code
        == 404
    )

    listed = viewer.get(f"/api/projects/{project['id']}/evidence").json()
    assert listed["total"] == 2 and listed["items"][0]["id"] == own["id"]
    assert viewer.delete(f"/api/evidence/{ev['id']}").status_code == 403
    assert annotator.delete(f"/api/evidence/{own['id']}").status_code == 204

    for fmt in ("docx", "xlsx", "csv"):
        job = viewer.post(
            f"/api/projects/{project['id']}/evidence/export", json={"format": fmt}
        ).json()
        drain()
        r = viewer.get(f"/api/jobs/{job['id']}/download")
        assert r.status_code == 200, r.text
        if fmt == "docx":
            doc = DocxDocument(io.BytesIO(r.content))
            text = [p.text for p in doc.paragraphs]
            assert text == [
                "证据篮",
                "1. Language is a system.",
                "出处：a · 句 1 · 作文",
                "查询式：language",
                "备注：定义句",
            ]
            assert [r.text for r in doc.paragraphs[1].runs if r.bold] == ["Language"]
            assert doc.styles["Normal"].font.name == "宋体"
        if fmt == "csv":
            rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
            assert rows == [
                ["原文", "命中", "译文", "出处", "查询式", "备注"],
                [
                    "Language is a system.",
                    "Language",
                    "",
                    "a · 句 1 · 作文",
                    "language",
                    "定义句",
                ],
            ]


def bearer_client(token):
    from services.api.main import app

    return TestClient(
        app, headers={"Authorization": f"Bearer {token}"}
    )  # no Origin, no cookie


def test_agent_token(owner, project, corpus, make_user, fake_stats):
    other = owner.post("/api/projects", json={"name": "另一个"}).json()
    r = owner.post("/api/auth/agent-token", json={"project_id": project["id"]})
    assert r.status_code == 200, r.text
    token = r.json()["token"]
    expires = datetime.fromisoformat(r.json()["expires_at"])
    assert timedelta(minutes=29) < expires - datetime.now(UTC) <= timedelta(minutes=30)

    agent = bearer_client(token)
    assert agent.get("/api/auth/me").json()["id"] == owner.user_id
    assert [p["id"] for p in agent.get("/api/projects").json()["items"]] == [
        project["id"]
    ]
    assert agent.get(f"/api/projects/{project['id']}").status_code == 200
    assert agent.get(f"/api/projects/{other['id']}").status_code == 403
    assert agent.get(f"/api/projects/{other['id']}/evidence").status_code == 403
    assert agent.post("/api/projects", json={"name": "x"}).status_code == 403
    assert (
        agent.post(
            "/api/auth/agent-token", json={"project_id": project["id"]}
        ).status_code
        == 403
    )

    # writes without Origin work with a bearer token; cookie requests still need Origin
    r = agent.post(
        f"/api/projects/{project['id']}/evidence",
        json={"kind": "text", "text": "agent 记录", "source": "助手"},
    )
    assert r.status_code == 200, r.text
    job = agent.post(
        "/api/stats/frequency", json={"scope": {"corpus_id": corpus["id"]}}
    ).json()
    assert agent.get(f"/api/jobs/{job['id']}").status_code == 200
    other_corpus = owner.post(
        f"/api/projects/{other['id']}/corpora", json={"name": "c", "language": "en"}
    ).json()
    assert agent.get(f"/api/corpora/{other_corpus['id']}").status_code == 403
    other_job = upload(owner, other_corpus["id"], {"x.txt": "hello"}).json()
    assert agent.get(f"/api/jobs/{other_job['id']}").status_code == 403

    assert bearer_client("nope").get("/api/auth/me").status_code == 401
    from sqlalchemy import update

    from services.api.db import transaction
    from services.api.models import AgentToken

    with transaction() as db:
        db.execute(
            update(AgentToken).values(
                expires_at=datetime.now(UTC) - timedelta(seconds=1)
            )
        )
    assert agent.get("/api/auth/me").status_code == 401

    outsider = make_user("x@example.edu")
    assert (
        outsider.post(
            "/api/auth/agent-token", json={"project_id": project["id"]}
        ).status_code
        == 404
    )


def test_pos_validated_against_corpus_tagset(owner, corpus, fake_stats, monkeypatch):
    from engine import cwb
    from services.api.storage import corpus_key

    scope = {"corpus_id": corpus["id"]}
    tagsets = {}
    monkeypatch.setattr(cwb, "tagset", lambda d, key: tagsets.get(key, "upos"))

    def code(path, body):
        r = owner.post(path, json=body)
        return r.json()["error"]["code"] if r.status_code == 422 else r.status_code

    assert (
        code("/api/stats/frequency", {"scope": scope, "pos": "nr"})
        == "VALIDATION_ERROR"
    )
    assert code("/api/search", {"scope": scope, "query": "语言/nr"}) == "QUERY_INVALID"
    tagsets[corpus_key(corpus["id"])] = "jieba"
    for pos in ("nr", "n", "NOUN", "uj"):
        assert code("/api/stats/frequency", {"scope": scope, "pos": pos}) == 200
    pattern = {"scope": scope, "n": 2, "pos_pattern": ["v", "DET"]}
    assert code("/api/stats/ngrams", pattern) == "VALIDATION_ERROR"
    assert code("/api/search", {"scope": scope, "query": "语言/nr"}) == 200
    assert code("/api/search", {"scope": scope, "query": "语言/zz"}) == "QUERY_INVALID"


def test_agent_token_cannot_delete_or_change_corpora(
    owner, project, corpus, fake_stats
):
    """Deletions, uploads and reindexing stay with the signed-in user (contracts/p2.md §九)."""
    pid, cid = project["id"], corpus["id"]
    token = owner.post("/api/auth/agent-token", json={"project_id": pid}).json()
    agent = bearer_client(token["token"])
    doc = owner.get(f"/api/corpora/{cid}/documents").json()["items"][0]
    sub = owner.post(
        f"/api/projects/{pid}/subcorpora",
        json={
            "name": "新闻",
            "corpus_id": cid,
            "filter": {"year": {"gte": 2000, "lte": 2005}},
        },
    ).json()
    saved = owner.post(
        f"/api/projects/{pid}/queries",
        json={"name": "q", "query": "language", "scope": {"corpus_id": cid}},
    ).json()
    stat_id = run_stat(owner, "frequency", {"scope": {"corpus_id": cid}})
    evidence = {"kind": "text", "text": "记录", "source": "笔记"}
    ev = owner.post(f"/api/projects/{pid}/evidence", json=evidence).json()
    own_ev = agent.post(f"/api/projects/{pid}/evidence", json=evidence).json()

    denied = {
        "corpus": agent.delete(f"/api/corpora/{cid}"),
        "document": agent.delete(f"/api/corpora/{cid}/documents/{doc['id']}"),
        "upload": upload(agent, cid, {"c.txt": "More language."}),
        "reindex": agent.post(f"/api/corpora/{cid}/reindex"),
        "subcorpus": agent.delete(f"/api/subcorpora/{sub['id']}"),
        "query": agent.delete(f"/api/queries/{saved['id']}"),
        "stat": agent.delete(f"/api/stats/{stat_id}"),
        "evidence": agent.delete(f"/api/evidence/{ev['id']}"),
        "own evidence": agent.delete(f"/api/evidence/{own_ev['id']}"),
    }
    assert {k: r.status_code for k, r in denied.items()} == dict.fromkeys(denied, 403)
    assert owner.get(f"/api/corpora/{cid}").json()["status"] == "ready"

    for path in (
        f"/api/evidence/{ev['id']}",
        f"/api/stats/{stat_id}",
        f"/api/queries/{saved['id']}",
        f"/api/subcorpora/{sub['id']}",
        f"/api/corpora/{cid}/documents/{doc['id']}",
    ):
        assert owner.delete(path).status_code == 204, path
    assert owner.post(f"/api/corpora/{cid}/reindex").status_code == 200
    assert owner.delete(f"/api/corpora/{cid}").status_code == 204


def test_job_creator_manages_own_job(owner, project, corpus, member, fake_stats):
    """Any member may cancel / retry the jobs they started; owner / editor manage all."""
    annotator, other, editor = (
        member("annotator"),
        member("viewer"),
        member("editor"),
    )
    body = {"scope": {"corpus_id": corpus["id"]}, "unit": "word"}
    job = annotator.post("/api/stats/frequency", json=body).json()
    assert job["owner_id"] == annotator.user_id
    assert job["owner_email"] == "annotator@example.edu"
    assert job["summary"] == "频次 · 作文"
    assert other.post(f"/api/jobs/{job['id']}/cancel").status_code == 403
    r = annotator.post(f"/api/jobs/{job['id']}/cancel")
    assert r.status_code == 200 and r.json()["state"] == "cancelled"
    retry = annotator.post(f"/api/jobs/{job['id']}/retry")
    assert retry.status_code == 200 and retry.json()["owner_id"] == annotator.user_id
    assert other.post(f"/api/jobs/{retry.json()['id']}/cancel").status_code == 403
    r = editor.post(f"/api/jobs/{retry.json()['id']}/cancel")
    assert r.status_code == 200 and r.json()["state"] == "cancelled"

    listed = owner.get(f"/api/projects/{project['id']}/jobs").json()["items"]
    stats_jobs = [j for j in listed if j["kind"] == "stats"]
    assert {j["owner_email"] for j in stats_jobs} == {"annotator@example.edu"}
    imports = [j for j in listed if j["kind"] == "import"]
    assert {j["summary"] for j in imports} == {"作文：a.txt", "作文：b.txt"}
    sub = owner.post(
        f"/api/projects/{project['id']}/subcorpora",
        json={"name": "新闻", "corpus_id": corpus["id"], "filter": {"genre": "news"}},
    ).json()
    stat_id = run_stat(
        owner,
        "keywords",
        {"target": body["scope"], "reference": {"subcorpus_id": sub["id"]}},
    )
    export = owner.post(f"/api/stats/{stat_id}/export", json={"format": "csv"}).json()
    assert export["summary"] == "关键词 · 作文 · CSV"


def test_evidence_same_hit_is_not_added_twice(owner, project, corpus):
    """Manual pin and the assistant's evidence.add of the same hit share one basket item."""
    h = owner.post(
        "/api/search", json={"scope": {"corpus_id": corpus["id"]}, "query": "language"}
    ).json()["items"][0]
    url = f"/api/projects/{project['id']}/evidence"
    manual = owner.post(
        url,
        json={
            "kind": "hit",
            "ref": {
                "hit_id": h["hit_id"],
                "document_id": h["document_id"],
                "start_cp": h["start_cp"],
                "end_cp": h["end_cp"],
            },
            "text": "KWIC",
            "source": "a",
        },
    ).json()
    assert manual["existing"] is False
    by_hit_id = owner.post(
        url,
        json={
            "kind": "hit",
            "ref": {"hit_id": h["hit_id"]},
            "text": "助手",
            "source": "a",
        },
    ).json()
    by_offsets = owner.post(
        url,
        json={
            "kind": "hit",
            "ref": {
                "document_id": h["document_id"],
                "start_cp": h["start_cp"],
                "end_cp": h["end_cp"],
            },
            "text": "助手",
            "source": "a",
        },
    ).json()
    assert by_hit_id["id"] == by_offsets["id"] == manual["id"]
    assert by_hit_id["existing"] is True and by_offsets["existing"] is True
    other = owner.post(
        url,
        json={
            "kind": "hit",
            "ref": {
                "document_id": h["document_id"],
                "start_cp": h["start_cp"] + 1,
                "end_cp": h["end_cp"],
            },
            "text": "另一条",
            "source": "a",
        },
    ).json()
    assert other["id"] != manual["id"] and other["ref"]["hit_id"]
    assert owner.get(url).json()["total"] == 2
