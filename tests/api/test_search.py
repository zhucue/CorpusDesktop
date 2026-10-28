import csv
import io
import json

import pytest
from openpyxl import load_workbook

from tests.api.conftest import drain, upload

TEXT_A = "Language is a system.\nThe language of science."
TEXT_B = "Grammar rules apply.\nA language again."


@pytest.fixture
def corpus(owner, project):
    c = owner.post(
        f"/api/projects/{project['id']}/corpora",
        json={"name": "作文", "language": "en"},
    ).json()
    upload(owner, c["id"], {"a.txt": TEXT_A}, metadata=json.dumps({"year": 2001}))
    upload(owner, c["id"], {"b.txt": TEXT_B}, metadata=json.dumps({"year": 2010}))
    drain()
    return c


def search(client, scope, query="language", **kw):
    return client.post("/api/search", json={"scope": scope, "query": query, **kw})


def test_parse_endpoint(owner):
    ok = owner.post("/api/search/parse", json={"query": "语言 +文化"}).json()
    assert ok["ok"] is True and ok["error"] is None
    assert ok["anchor"] == "语言"
    assert ok["conditions"] == [{"op": "+", "text": "文化"}]
    assert ok["tokens"]
    attr = owner.post("/api/search/parse", json={"query": "学生{dep=nsubj}"}).json()
    assert attr["ok"] and [t["kind"] for t in attr["tokens"]] == ["word", "attr"]
    bad = owner.post("/api/search/parse", json={"query": "{dep=nsubj"}).json()
    assert (
        bad["ok"] is False and bad["error"]["position"] == 0 and bad["error"]["message"]
    )


def test_search_hits_counts_and_paging(owner, corpus, fake_engine, member):
    r = search(member("viewer"), {"corpus_id": corpus["id"]}, page_size=2)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["counts"] == {
        "hits": 3,
        "sentences": 2 + 1,
        "documents": 2,
        "complete": True,
    }
    assert body["parse"]["ok"] and body["page"] == 1 and body["page_size"] == 2
    hit = body["items"][1]
    assert hit["document_title"] == "a" and hit["metadata"] == {"year": 2001}
    assert (
        hit["match"] == "language"
        and hit["sentence_index"] == 1
        and hit["aligned"] == []
    )
    assert hit["hit_id"] == f"{hit['document_id']}:{hit['start_cp']}-{hit['end_cp']}"

    call = fake_engine.calls[-1]
    assert call["restrict"] is None and call["page_size"] == 2
    assert call["context"].unit == "char" and call["context"].size == 40

    page2 = search(
        owner,
        {"corpus_id": corpus["id"]},
        page=2,
        page_size=2,
        search_id=body["search_id"],
    ).json()
    assert page2["search_id"] == body["search_id"] and len(page2["items"]) == 1

    # a search_id only matches the same query and scope
    r = search(
        owner, {"corpus_id": corpus["id"]}, query="grammar", search_id=body["search_id"]
    )
    assert r.status_code == 409 and r.json()["error"]["code"] == "SEARCH_EXPIRED"


def test_incomplete_counts_then_recheck(owner, corpus, fake_engine):
    fake_engine.complete = False
    first = search(owner, {"corpus_id": corpus["id"]}).json()
    assert first["counts"] == {
        "hits": None,
        "sentences": None,
        "documents": None,
        "complete": False,
    }
    fake_engine.complete = True
    again = search(
        owner, {"corpus_id": corpus["id"]}, search_id=first["search_id"]
    ).json()
    assert again["counts"]["complete"] and again["counts"]["hits"] == 3


def test_random_sort_pins_seed(owner, corpus, fake_engine):
    first = search(owner, {"corpus_id": corpus["id"]}, sort="random").json()
    seed = fake_engine.calls[-1]["seed"]
    assert seed is not None
    search(
        owner,
        {"corpus_id": corpus["id"]},
        sort="random",
        page=2,
        search_id=first["search_id"],
    )
    assert fake_engine.calls[-1]["seed"] == seed


def test_search_errors(owner, project, corpus, make_user):
    r = search(owner, {"corpus_id": corpus["id"]}, query="{dep=nsubj")
    assert r.status_code == 422 and r.json()["error"]["code"] == "QUERY_INVALID"
    empty = owner.post(
        f"/api/projects/{project['id']}/corpora", json={"name": "空", "language": "en"}
    ).json()
    r = search(owner, {"corpus_id": empty["id"]})
    assert r.status_code == 409 and r.json()["error"]["code"] == "CORPUS_NOT_READY"
    assert (
        search(make_user("x@example.edu"), {"corpus_id": corpus["id"]}).status_code
        == 404
    )
    r = owner.post("/api/search", json={"scope": {}, "query": "x"})
    assert r.status_code == 422


def test_subcorpus_scope_passes_restrict(owner, project, corpus, fake_engine):
    sub = owner.post(
        f"/api/projects/{project['id']}/subcorpora",
        json={
            "name": "2001",
            "corpus_id": corpus["id"],
            "filter": {"year": {"gte": 2000, "lte": 2005}},
        },
    ).json()
    r = search(owner, {"subcorpus_id": sub["id"]})
    assert r.status_code == 200, r.text
    restrict = fake_engine.calls[-1]["restrict"]
    assert restrict and "year" in restrict


def test_parallel_hits_carry_aligned_sentences(owner, project, fake_engine):
    c = owner.post(
        f"/api/projects/{project['id']}/corpora",
        json={"name": "双语", "language": "en"},
    ).json()
    upload(
        owner, c["id"], {"pair.tmx": "I study language.\t我研究语言。\nHello.\t你好。"}
    )
    drain()
    corpus = owner.get(f"/api/corpora/{c['id']}").json()
    assert corpus["parallel"] is True and corpus["status"] == "ready"
    hit = search(owner, {"corpus_id": c["id"]}).json()["items"][0]
    assert hit["aligned"] == [{"language": "zh", "text": "我研究语言。"}]


def test_context(owner, corpus):
    hit = search(owner, {"corpus_id": corpus["id"]}).json()["items"][1]
    r = owner.get(
        "/api/search/context",
        params={
            "document_id": hit["document_id"],
            "sentence_index": hit["sentence_index"],
            "start_cp": hit["start_cp"],
            "end_cp": hit["end_cp"],
        },
    ).json()
    assert r["document"]["title"] == "a"
    assert [s["index"] for s in r["sentences"]] == [0, 1]
    target = r["sentences"][1]
    a, b = target["highlight"]
    assert target["text"][a:b] == "language"
    assert r["sentences"][0]["highlight"] is None


def test_export_csv_and_xlsx(owner, project, corpus, member):
    viewer = member("viewer")
    job = viewer.post(
        "/api/search/export",
        json={
            "scope": {"corpus_id": corpus["id"]},
            "query": "language",
            "format": "csv",
        },
    ).json()
    assert job["kind"] == "export" and job["project_id"] == project["id"]
    assert viewer.get(f"/api/jobs/{job['id']}/download").status_code == 409
    drain()
    done = viewer.get(f"/api/jobs/{job['id']}").json()
    assert done["state"] == "succeeded", done
    assert done["result"]["rows"] == done["result"]["hits"] == 3
    r = viewer.get(f"/api/jobs/{job['id']}/download")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][:3] == ["命中 ID", "文档", "文档 ID"] and rows[0][-1] == "year"
    assert len(rows) - 1 == 3

    job = owner.post(
        "/api/search/export",
        json={
            "scope": {"corpus_id": corpus["id"]},
            "query": "language",
            "format": "xlsx",
        },
    ).json()
    drain()
    r = owner.get(f"/api/jobs/{job['id']}/download")
    ws = load_workbook(io.BytesIO(r.content)).active
    header = next(ws.iter_rows(min_row=1, max_row=1))
    assert all(c.font.name == "宋体" for c in header)
    assert all(c.fill.fill_type is None for row in ws.iter_rows() for c in row)
    assert ws.max_row == 4


def test_export_builtin_corpus_into_project(owner, project, make_user):
    admin = make_user("admin@example.edu", admin=True)
    builtin = admin.post(
        "/api/corpora", json={"name": "UD 英文", "language": "en"}
    ).json()
    upload(admin, builtin["id"], {"ud.txt": TEXT_A})
    drain()
    body = {"scope": {"corpus_id": builtin["id"]}, "query": "language", "format": "csv"}
    job = owner.post(
        "/api/search/export", json={**body, "project_id": project["id"]}
    ).json()
    assert job["project_id"] == project["id"]
    outsider = make_user("x@example.edu")
    assert (
        outsider.post(
            "/api/search/export", json={**body, "project_id": project["id"]}
        ).status_code
        == 404
    )


def test_saved_queries(owner, project, corpus, member):
    pid = project["id"]
    body = {
        "name": "language",
        "query": "language",
        "scope": {"corpus_id": corpus["id"]},
    }
    assert (
        member("viewer").post(f"/api/projects/{pid}/queries", json=body).status_code
        == 403
    )
    saved = owner.post(f"/api/projects/{pid}/queries", json=body).json()
    assert (
        saved["scope"] == {"corpus_id": corpus["id"]} and saved["query"] == "language"
    )
    assert owner.get(f"/api/projects/{pid}/queries").json()["total"] == 1
    assert owner.delete(f"/api/queries/{saved['id']}").status_code == 204
    assert owner.get(f"/api/projects/{pid}/queries").json()["total"] == 0


def test_include_tokens(owner, corpus):
    plain = search(owner, {"corpus_id": corpus["id"]}).json()
    assert "tokens" not in plain["items"][0]
    items = search(owner, {"corpus_id": corpus["id"]}, include_tokens=True).json()[
        "items"
    ]
    assert items[0]["tokens"] == [
        {"part": "match", "form": "Language", "pos": None},
        {"part": "right", "form": "is", "pos": None},
        {"part": "right", "form": "a", "pos": None},
        {"part": "right", "form": "system.", "pos": None},
    ]
    assert [(t["part"], t["form"]) for t in items[1]["tokens"]] == [
        ("left", "The"),
        ("match", "language"),
        ("right", "of"),
        ("right", "science."),
    ]


def test_parallel_context_and_evidence_stay_in_sentence(owner, project, fake_engine):
    """QA Q20/Q60: parallel corpora search with sentence-only context; a pinned hit keeps its own
    sentence, the aligned translation and the sentence number, with the match in bold."""
    c = owner.post(
        f"/api/projects/{project['id']}/corpora",
        json={"name": "双语", "language": "en"},
    ).json()
    upload(
        owner,
        c["id"],
        {"pair.tmx": "Hello.\t你好。\nI study language.\t我研究语言。\nBye.\t再见。"},
    )
    drain()
    hit = search(owner, {"corpus_id": c["id"]}).json()["items"][0]
    assert fake_engine.calls[-1]["context"].sentence_only is True
    ref = {
        "hit_id": hit["hit_id"],
        "document_id": hit["document_id"],
        "start_cp": hit["start_cp"],
        "end_cp": hit["end_cp"],
        "search_query": "language",
    }
    ev = owner.post(
        f"/api/projects/{project['id']}/evidence",
        json={"kind": "hit", "ref": ref, "text": "x", "source": "y"},
    ).json()
    assert ev["text"] == "I study **language**.\nZH：我研究语言。"
    assert ev["source"] == "pair-en · 句 2 · 双语"
    assert ev["ref"]["sentence"] == "I study language." and ev["ref"]["match"] == [
        8,
        16,
    ]
    assert ev["ref"]["aligned"] == [{"language": "zh", "text": "我研究语言。"}]


def test_parse_with_scope_checks_layers(owner, corpus, monkeypatch):
    """QA Q21/Q22: parse against the corpus (tagset, POS tags) and report a missing layer before searching."""
    from engine import cwb

    scope = {"corpus_id": corpus["id"]}
    r = owner.post("/api/search/parse", json={"query": "language/n", "scope": scope})
    body = r.json()
    assert body["ok"] and body["corpus"]["tagset"] in ("upos", "jieba", "mixed")
    assert body["corpus"]["layers"] == {"dep": False, "span": False}
    monkeypatch.setattr(cwb, "has_values", lambda ix, key, attr: False)
    bad = owner.post(
        "/api/search/parse", json={"query": "language{dep=nsubj}", "scope": scope}
    ).json()
    assert bad["ok"] is False and bad["error"]["position"] == 8
    assert "依存" in bad["error"]["message"]
    span = owner.post(
        "/api/search/parse", json={"query": "[时态]", "scope": scope}
    ).json()
    assert span["ok"] is False and "片段" in span["error"]["message"]
    r = search(owner, scope, "{dep=nsubj}")
    assert r.status_code == 422 and r.json()["error"]["code"] == "INVALID_LAYER"
