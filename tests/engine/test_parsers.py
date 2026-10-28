from pathlib import Path

import pytest

from engine.model import Alignment
from engine.parsers import ParseError, decode, parse, table

FIX = Path(__file__).resolve().parents[1] / "fixtures"
IMP = FIX / "import"


def check_offsets(doc):
    for s in doc.sentences:
        assert 0 <= s.start <= s.end <= len(doc.text)
        for t in s.tokens:
            assert s.start <= t.start <= t.end <= s.end


def test_txt():
    [doc], al = parse(FIX / "sample.txt", language="zh", metadata={"genre": "test"})
    assert (
        al == []
        and doc.title == "sample"
        and doc.language == "zh"
        and doc.metadata == {"genre": "test"}
    )
    assert doc.text.split("\n") == [
        "语言是人类最重要的交际工具。社会语言学研究语言与社会的关系。",
        "第二段：语言文化与翻译。",
        "The last paragraph is English.",
    ]
    assert doc.sentences == []


def test_docx():
    [doc], _ = parse(FIX / "sample.docx", language=None, metadata={})
    assert doc.title == "示例文档" and doc.language == ""
    assert doc.text == "第一段：语言与翻译。\nSecond paragraph in English."


def test_csv_and_tsv():
    docs, _ = parse(FIX / "sample.csv", language="en", metadata={"batch": 1})
    assert [d.title for d in docs] == [
        "新闻一",
        "新闻二",
        "Essay",
    ]  # empty text row skipped
    assert docs[0].metadata == {
        "batch": 1,
        "year": 1990,
        "genre": "news",
        "code": "007",
    }
    assert (
        docs[1].text == "翻译研究关注文化差异，也关注语言。"
        and docs[2].language == "en"
    )
    docs, _ = parse(FIX / "sample.tsv", language="en", metadata={})
    assert [(d.title, d.text, d.metadata) for d in docs] == [
        ("Row A", "first row text", {"source": "web"}),
        ("sample #2", "second row", {"source": "book"}),
    ]


def test_csv_without_text_column_detects_defaults():
    docs, _ = parse(IMP / "headlines.csv", language="en", metadata={})
    # abstract is a text-like name, headline a title-like one; BOM, CRLF and quoted newlines handled
    assert [(d.title, d.text) for d in docs] == [
        ("50 States, 50 Fixes", "A series on climate initiatives."),
        ("Heat Waves", "Cities prepare for heat.\nA second line."),
    ]
    assert docs[0].metadata == {
        "record_key": "k1",
        "pub_date": "2026-01-02",
        "source": "nyt",
    }


def test_csv_mapping_joins_two_columns():
    mapping = {
        "text": ["headline", "abstract"],
        "title": "record_key",
        "language": None,
        "metadata": ["pub_date"],
    }
    docs, _ = parse(
        IMP / "headlines.csv", language="en", metadata={"b": 1}, mapping=mapping
    )
    assert [(d.title, d.text) for d in docs] == [
        ("k1", "50 States, 50 Fixes\n\nA series on climate initiatives."),
        ("k2", "Heat Waves\n\nCities prepare for heat.\nA second line."),
        ("k3", "No Abstract"),
    ]
    assert docs[0].metadata == {"b": 1, "pub_date": "2026-01-02"}


def test_csv_mapping_errors(tmp_path):
    with pytest.raises(ParseError, match="headlines.csv 缺少列：body"):
        parse(
            IMP / "headlines.csv",
            language="en",
            metadata={},
            mapping={"text": ["body"]},
        )
    with pytest.raises(ParseError, match="1–2 个正文列"):
        parse(IMP / "headlines.csv", language="en", metadata={}, mapping={"text": []})
    p = tmp_path / "x.csv"
    p.write_text("title,body\na,\n", encoding="utf-8")
    with pytest.raises(ParseError, match="x.csv 正文列"):
        parse(p, language="en", metadata={}, mapping={"text": ["body"]})


def test_csv_longest_column_and_preview(tmp_path):
    p = tmp_path / "y.csv"
    p.write_text(
        "id,a,b\n1,short,a much longer value here\n2,x,another long value\n",
        encoding="utf-8",
    )
    [d1, _], _ = parse(p, language="en", metadata={})
    assert d1.text == "a much longer value here"
    assert d1.metadata == {"id": 1, "a": "short"}
    data = (IMP / "headlines.csv").read_bytes()
    pv = table.preview(
        data[:150], "headlines.csv"
    )  # cut mid-file: partial line dropped
    assert pv["columns"] == ["record_key", "headline", "abstract", "pub_date", "source"]
    assert pv["rows"][0][:2] == ["k1", "50 States, 50 Fixes"]
    assert pv["defaults"] == {
        "text": ["abstract"],
        "title": "headline",
        "language": None,
        "metadata": None,
    }


def test_txt_encodings():
    [doc], _ = parse(IMP / "gbk.txt", language="zh", metadata={})
    assert doc.text == "气候变化影响全球。\n翻译研究关注语言与文化。"
    [doc], _ = parse(IMP / "utf16.txt", language="zh", metadata={})
    assert doc.text == "语料库语言学\nCorpus linguistics"
    assert decode("abc".encode("utf-16-le"), "x.txt") == "abc"
    assert decode(b"a\r\nb\rc", "x.txt") == "a\nb\nc"
    assert decode("中文".encode()[:-1], "x.txt", partial=True) == "中"


def test_docx_tables_headings_images():
    [doc], _ = parse(IMP / "table.docx", language="zh", metadata={})
    assert doc.title == "气候报告"  # first heading; no Title paragraph, no core title
    assert doc.text == "气候报告\n第一段正文。\n年份\n排放\n合并单元格\n结尾段落。"


def test_docx_title_style_wins(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("第一章", level=1)
    d.add_paragraph("翻译策略研究", style="Title")
    d.add_paragraph("正文。")
    d.core_properties.title = "模板标题"
    d.save(tmp_path / "file-name.docx")
    [doc], _ = parse(tmp_path / "file-name.docx", language="zh", metadata={})
    assert doc.title == "翻译策略研究"
    d.paragraphs[1].style = "Normal"
    d.save(tmp_path / "file-name.docx")
    [doc], _ = parse(tmp_path / "file-name.docx", language="zh", metadata={})
    assert doc.title == "模板标题"


def test_tmx_lang_attribute_and_inline_codes():
    docs, al = parse(IMP / "inline.tmx", language=None, metadata={})
    assert {d.language: d.text for d in docs} == {
        "en": "Press Save now.\nSecond unit.",
        "zh": "立即保存。\n第二个单元。",
    }
    assert al == [
        Alignment("inline:1:en", "inline:1:zh"),
        Alignment("inline:2:en", "inline:2:zh"),
    ]


def test_conllu_multiword_and_empty_nodes():
    [doc], _ = parse(IMP / "mwt.conllu", language="es", metadata={})
    [s] = doc.sentences
    assert doc.text == "Vámonos al mar." and s.external_id == "es-1"
    assert [(t.form, doc.text[t.start : t.end]) for t in s.tokens] == [
        ("Vamos", "Vámonos"),
        ("nos", "Vámonos"),
        ("a", "al"),
        ("el", "al"),
        ("mar", "mar"),
        (".", "."),
    ]
    check_offsets(doc)


def test_conllu_ud_english():
    docs, _ = parse(
        FIX / "ud-en-ewt-sample.conllu", language="en", metadata={"src": "ud"}
    )
    assert len(docs) == 16 and sum(len(d.sentences) for d in docs) == 221
    assert (
        docs[0].external_id
        == docs[0].title
        == "weblog-blogspot.com_zentelligence_20040423000200_ENG_20040423_000200"
    )
    s = docs[0].sentences[0]
    assert (
        s.external_id.endswith("-0001")
        and docs[0].text[s.start : s.end] == "What if Google Morphed Into GoogleOS?"
    )
    t = s.tokens[3]
    assert (t.form, t.pos, t.lemma, t.head, t.deprel) == (
        "Morphed",
        "VERB",
        "morph",
        1,
        "advcl",
    )
    for d in docs:
        check_offsets(d)
        for s in d.sentences:
            for t in s.tokens:
                assert (
                    d.text[t.start : t.end] == t.form
                )  # EWT forms are substrings of # text, incl. multiword parts
    toks = [
        t
        for d in docs
        for s in d.sentences
        for t in s.tokens
        if t.form in ("Google", "'s")
    ]
    assert toks  # "Google's" multiword token split into parts


def test_conllu_without_text(tmp_path):
    p = tmp_path / "x.conllu"
    p.write_text(
        "1\tHello\thello\tINTJ\t_\t_\t0\troot\t_\tSpaceAfter=No\n2\t!\t!\tPUNCT\t_\t_\t1\tpunct\t_\t_\n\n"
        "1-2\tdon't\t_\t_\t_\t_\t_\t_\t_\t_\n1\tdo\tdo\tAUX\t_\t_\t0\troot\t_\t_\n2\tn't\tnot\tPART\t_\t_\t1\tadvmod\t_\t_\n"
        "2.1\tx\tx\tX\t_\t_\t_\t_\t_\t_\n\n",
        encoding="utf-8",
    )
    [doc], _ = parse(p, language="en", metadata={})
    assert doc.text == "Hello! don't"
    assert [
        (doc.text[t.start : t.end], t.form) for s in doc.sentences for t in s.tokens
    ] == [("Hello", "Hello"), ("!", "!"), ("do", "do"), ("n't", "n't")]


def test_conllu_chinese_paragraphs():
    docs, _ = parse(FIX / "zh-sample.conllu", language="zh", metadata={})
    assert [d.title for d in docs] == ["zh1", "zh2"]
    assert docs[1].text == "他说语言。\n汉语言文字学"
    for d in docs:
        check_offsets(d)


def test_tmx():
    docs, al = parse(FIX / "sample.tmx", language=None, metadata={"set": "t"})
    by = {d.language: d for d in docs}
    assert set(by) == {"en", "zh"}
    assert (
        by["en"].text
        == "Language shapes thought.\nTranslation is a bridge.\nOnly English here."
    )
    assert by["zh"].text == "语言塑造思维。\n翻译是桥梁。"
    assert al == [
        Alignment("sample:1:en", "sample:1:zh"),
        Alignment("sample:2:en", "sample:2:zh"),
    ]
    ids = {s.external_id: d.text[s.start : s.end] for d in docs for s in d.sentences}
    assert (
        ids["sample:2:zh"] == "翻译是桥梁。"
        and ids["sample:3:en"] == "Only English here."
    )


def test_tatoeba_dir():
    docs, al = parse(
        FIX / "tatoeba", language=None, metadata={"license": "CC BY 2.0 FR"}
    )
    by = {d.language: d for d in docs}
    assert set(by) == {"zh", "en"}
    assert [by["zh"].text[s.start : s.end] for s in by["zh"].sentences] == [
        "我们试试看！",
        "我该去睡觉了。",
        "你在干什么啊？",
    ]
    assert [s.external_id for s in by["en"].sentences] == ["1276", "1277", "1280"]
    assert by["en"].metadata == {"license": "CC BY 2.0 FR", "tatoeba_lang": "eng"}
    # reciprocal 1/1276 collapses; 3-9999 has a missing endpoint; 1277-1280 is same-language
    assert al == [Alignment("1", "1276"), Alignment("2", "1277")]


def test_tatoeba_single_file():
    docs, al = parse(FIX / "tatoeba" / "cmn_sentences.tsv", language=None, metadata={})
    assert (
        [d.language for d in docs] == ["zh"]
        and len(docs[0].sentences) == 3
        and al == []
    )


def test_unknown_format(tmp_path):
    p = tmp_path / "x.pdf"
    p.write_bytes(b"%PDF")
    with pytest.raises(ParseError, match="不支持"):
        parse(p, language=None, metadata={})
