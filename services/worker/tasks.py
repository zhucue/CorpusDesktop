"""Job handlers. Each handler ends with one transaction that locks the job row,
writes results and marks the job succeeded, so a cancel either wins entirely or not at all."""

import csv
import json
import shutil
import time
import uuid
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from sqlalchemy import bindparam, func, insert, select, update
from sqlalchemy.dialects.postgresql import insert as pg_insert

from engine import cwb, nlp, parsers, query, stats
from engine import model as em
from engine.conllu import write_conllu
from services.api import effective
from services.api.annotations import (
    is_zh,
    refresh_progress,
)
from services.api.annotators import layer_annotators
from services.api.db import session, transaction
from services.api.jobs import corpus_tokens, enqueue, live_index
from services.api.metadata import field_types
from services.api.models import (
    Annotation,
    AnnotationLayer,
    Corpus,
    Document,
    Evidence,
    Job,
    Sentence,
    SentenceAlignment,
    StatResult,
    Subcorpus,
    Token,
    User,
)
from services.api.search import (
    ExportInput,
    LayersInput,
    Scope,
    corpus_tagset,
    enrich,
    plan_query,
    pure_attr,
    resolve_scope,
    scope_filter,
    search_page,
    with_span_layer,
)
from services.api.storage import (
    corpus_key,
    export_path,
    index_dir,
    stats_path,
    text_path,
)
from services.worker import annot_tasks, kb_tasks
from services.worker.main import Cancelled, JobFailed, log

EXPORT_PAGE = 1000
EXPORT_COUNT_TIMEOUT = 600
LOAD_BATCH = 500
# token rows read per index load round; a larger document loads alone
LOAD_TOKENS = 1_000_000
WRITE_SENTENCES = 5000  # sentences (with their tokens) per INSERT round at import
PROGRESS_SECONDS = 2.0


def _corpus(db, corpus_id):
    corpus = db.get(Corpus, corpus_id)
    if not corpus:
        raise JobFailed("CORPUS_DELETED", "语料库已删除")
    return corpus


# ---- import ----


def _parse_batch(run):
    params = run.params
    folder = Path(params["dir"])
    names = params["files"]
    # Tatoeba sentences + links files are parsed together as one directory.
    targets = (
        [folder]
        if any("links" in n.lower() for n in names)
        else [folder / n for n in names]
    )
    documents, alignments = [], []
    for i, path in enumerate(targets):
        run.progress(i, len(targets), f"解析 {path.name}", "个文件")
        try:
            docs, links = parsers.parse(
                path,
                language=params["language"],
                metadata=params["metadata"],
                mapping=params.get("mapping"),
            )
        except parsers.ParseError as e:
            raise JobFailed("PARSE_ERROR", str(e)) from e
        except Exception as e:  # a malformed file must not surface as INTERNAL
            log.exception("parsing %s failed", path)
            raise JobFailed("PARSE_ERROR", f"{path.name} 无法解析：{e}"[:500]) from e
        documents += docs
        alignments += links
    if not documents:
        raise JobFailed("EMPTY_IMPORT", "文件中没有可导入的文本")
    return documents, alignments


def _write_documents(db, corpus_id, documents, written):
    """Insert documents/sentences/tokens; returns sentence ids keyed by external id."""
    by_external = {}
    for doc in documents:
        doc_id = uuid.uuid4()
        path = text_path(corpus_id, doc_id)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(doc.text, encoding="utf-8", newline="")
        written.append(path)
        db.execute(
            insert(Document),
            {
                "id": doc_id,
                "corpus_id": corpus_id,
                "title": doc.title,
                "language": doc.language,
                "meta": doc.metadata,
                "external_id": doc.external_id,
                "text_path": str(path),
                "char_count": len(doc.text),
                "sentence_count": len(doc.sentences),
                "token_count": sum(len(s.tokens) for s in doc.sentences),
            },
        )
        for k in range(0, len(doc.sentences), WRITE_SENTENCES):
            part = doc.sentences[k : k + WRITE_SENTENCES]
            ids = db.scalars(
                insert(Sentence).returning(Sentence.id, sort_by_parameter_order=True),
                [
                    {
                        "document_id": doc_id,
                        "idx": k + i,
                        "start_cp": s.start,
                        "end_cp": s.end,
                        "text": doc.text[s.start : s.end],
                    }
                    for i, s in enumerate(part)
                ],
            ).all()
            tokens = [
                {
                    "sentence_id": sid,
                    "idx": j,
                    "start_cp": t.start,
                    "end_cp": t.end,
                    "form": t.form,
                    "pos": t.pos,
                    "lemma": t.lemma,
                    "head": t.head,
                    "deprel": t.deprel,
                }
                for sid, s in zip(ids, part)
                for j, t in enumerate(s.tokens)
            ]
            if tokens:
                db.execute(insert(Token), tokens)
            for sid, s in zip(ids, part):
                if s.external_id:
                    by_external[s.external_id] = sid
    return by_external


def run_import(run):
    corpus_id = uuid.UUID(run.params["corpus_id"])
    documents, alignments = _parse_batch(run)
    total = len(documents)
    # progress in characters so one large document (Tatoeba: one per language) still advances
    chars = sum(len(d.text) for d in documents)
    done, last = 0, 0.0

    def report(offset):
        nonlocal last
        if time.monotonic() - last >= PROGRESS_SECONDS:
            last = time.monotonic()
            run.progress(done + offset, chars, "分句与标注", "字")

    # language "auto" (or a CSV row without one): detect from the text
    for d in documents:
        if not d.language:
            d.language = nlp.detect_language(d.text)

    # CoNLL-U arrives tokenized and is stored as is, whatever its language
    todo = [
        d for d in documents if not (d.sentences and all(s.tokens for s in d.sentences))
    ]
    done = chars - sum(len(d.text) for d in todo)
    report(0)
    try:  # Stanza batches paragraphs across documents (many short CSV rows)
        nlp.annotate_many(todo, report)
    except ValueError as e:  # e.g. a language without an annotation model
        raise JobFailed("ANNOTATE_ERROR", str(e)[:500]) from e
    run.progress(chars, chars, "写入", "字")
    written = []
    try:
        with transaction() as db:
            run.lock(db)
            corpus = _corpus(db, corpus_id)
            by_external = _write_documents(db, corpus_id, documents, written)
            pairs = {
                (by_external[a.src_external_id], by_external[a.tgt_external_id])
                for a in alignments
                if a.src_external_id in by_external and a.tgt_external_id in by_external
            }
            if pairs:
                db.execute(
                    insert(SentenceAlignment),
                    [{"src_sentence_id": s, "tgt_sentence_id": t} for s, t in pairs],
                )
                corpus.parallel = True
            if corpus.language in ("auto", ""):
                shares = Counter()
                for d in documents:
                    shares[nlp.lang_code(d.language)] += len(d.text)
                corpus.language = shares.most_common(1)[0][0] or corpus.language
            index_job = db.scalar(
                select(Job).where(
                    Job.kind == "index",
                    Job.state == "queued",
                    Job.params["corpus_id"].astext == str(corpus_id),
                )
            )
            tokens = corpus_tokens(db, corpus_id)
            if index_job is None:
                enqueue(
                    db,
                    "index",
                    {"corpus_id": str(corpus_id), "tokens": tokens},
                    run.owner_id,
                    run.project_id,
                )
            else:
                index_job.params = {**index_job.params, "tokens": tokens}
            # a searchable corpus keeps serving its old index until the new one is in
            if not live_index(corpus):
                corpus.status = "indexing"
            corpus.error = None
            run.succeed(
                db,
                {
                    "documents": total,
                    "sentences": sum(len(d.sentences) for d in documents),
                    "tokens": sum(
                        len(s.tokens) for d in documents for s in d.sentences
                    ),
                    "alignments": len(pairs),
                },
            )
    except BaseException:
        for path in written:
            path.unlink(missing_ok=True)
        raise
    shutil.rmtree(run.params["dir"], ignore_errors=True)


# ---- index ----


def _load_batches(docs):
    """Group (id, token_count) rows into load rounds of at most LOAD_BATCH documents and LOAD_TOKENS
    tokens; a document larger than LOAD_TOKENS is a round of its own."""
    batch, tokens = [], 0
    for doc_id, n in docs:
        if batch and (len(batch) >= LOAD_BATCH or tokens + n > LOAD_TOKENS):
            yield batch
            batch, tokens = [], 0
        batch.append(doc_id)
        tokens += n
    if batch:
        yield batch


def _load_documents(docs, run):
    """Stream engine Documents from the database in the given order, one load round at a time.
    Rows are read as plain tuples; ORM objects for millions of tokens would not fit in memory."""
    total = len(docs)
    done = 0
    for batch in _load_batches(docs):
        run.progress(done, total, "建索引", "篇")
        with session() as db:
            rows = {
                r.id: r
                for r in db.execute(
                    select(
                        Document.id,
                        Document.title,
                        Document.language,
                        Document.meta,
                        Document.external_id,
                        Document.text_path,
                    ).where(Document.id.in_(batch))
                )
            }
            sentences = db.execute(
                select(
                    Sentence.id,
                    Sentence.document_id,
                    Sentence.start_cp,
                    Sentence.end_cp,
                )
                .where(Sentence.document_id.in_(batch))
                .order_by(Sentence.document_id, Sentence.idx)
            ).all()
            by_sentence = {}
            for sid, *fields in db.execute(
                select(
                    Token.sentence_id,
                    Token.start_cp,
                    Token.end_cp,
                    Token.form,
                    Token.pos,
                    Token.lemma,
                    Token.head,
                    Token.deprel,
                )
                .join(Sentence, Sentence.id == Token.sentence_id)
                .where(Sentence.document_id.in_(batch))
                .order_by(Token.sentence_id, Token.idx)
                .execution_options(yield_per=50_000)
            ):
                by_sentence.setdefault(sid, []).append(em.Token(*fields))
        by_doc = {}
        for sid, doc_id, a, b in sentences:
            by_doc.setdefault(doc_id, []).append(
                em.Sentence(a, b, by_sentence.pop(sid, []))
            )
        del sentences
        for doc_id in batch:
            d = rows.get(doc_id)
            if d is None:
                raise JobFailed("DOCUMENTS_CHANGED", "建索引期间文档被删除，请重试")
            yield em.Document(
                title=d.title,
                language=d.language,
                text=Path(d.text_path).read_text(encoding="utf-8"),
                metadata=d.meta,
                sentences=by_doc.pop(doc_id, []),
                external_id=d.external_id,
            )
        done += len(batch)


def run_index(run):
    corpus_id = uuid.UUID(run.params["corpus_id"])
    with transaction() as db:
        corpus = _corpus(db, corpus_id)
        if not live_index(corpus):
            corpus.status = "indexing"
        docs = db.execute(
            select(Document.id, Document.token_count)
            .where(Document.corpus_id == corpus_id)
            .order_by(Document.created_at, Document.id)
        ).all()
    info = None
    if docs:
        info = cwb.build(
            index_dir(corpus_id),
            corpus_key(corpus_id),
            _load_documents(docs, run),
            [str(i) for i, _ in docs],
        )
    with transaction() as db:
        run.lock(db)
        corpus = _corpus(db, corpus_id)
        docs, sentences, tokens = db.execute(
            select(
                func.count(),
                func.coalesce(func.sum(Document.sentence_count), 0),
                func.coalesce(func.sum(Document.token_count), 0),
            ).where(Document.corpus_id == corpus_id)
        ).one()
        stats = {"documents": docs, "sentences": int(sentences), "tokens": int(tokens)}
        if corpus.parallel:
            stats["pairs"] = _pair_count(db, corpus_id)
        corpus.stats = stats
        corpus.meta = corpus_profile(db, corpus, info.tagset if info else None)
        corpus.index_version = info.index_version if info else None
        corpus.status = "ready" if info else "empty"
        corpus.error = None
        run.succeed(db, {**corpus.stats, "index_version": corpus.index_version})


def _pair_count(db, corpus_id):
    """Sentence pairs whose source sentence is in the corpus (parallel corpora)."""
    return db.scalar(
        select(func.count())
        .select_from(SentenceAlignment)
        .join(Sentence, Sentence.id == SentenceAlignment.src_sentence_id)
        .join(Document, Document.id == Sentence.document_id)
        .where(Document.corpus_id == corpus_id)
    )


PROFILE_SENTENCES = 3000


def corpus_profile(db, corpus, tagset):
    """corpus.meta with the index profile filled in: tagset (upos | jieba | mixed, as the query engine
    resolves POS), languages (document language codes by text share), script of the Chinese text
    (simplified | traditional | mixed; absent without enough Chinese) and indexed_at."""
    meta = {
        k: v
        for k, v in (corpus.meta or {}).items()
        if k not in ("tagset", "script", "languages", "indexed_at")
    }
    shares, raw_zh = Counter(), []
    for raw, chars in db.execute(
        select(Document.language, func.sum(Document.char_count))
        .where(Document.corpus_id == corpus.id)
        .group_by(Document.language)
    ):
        code = nlp.lang_code(raw)
        shares[code] += int(chars or 0)
        if code == "zh":
            raw_zh.append(raw)
    if not shares:
        return meta
    meta["languages"] = [c for c, _ in shares.most_common() if c]
    if tagset:
        meta["tagset"] = tagset
    if raw_zh:
        # spread the sample over the whole corpus: the first sentences can be all one script (Tatoeba)
        zh = (
            select(Sentence.id, Sentence.text)
            .join(Document, Document.id == Sentence.document_id)
            .where(Document.corpus_id == corpus.id, Document.language.in_(raw_zh))
            .subquery()
        )
        n = db.scalar(select(func.count()).select_from(zh)) or 0
        step = max(1, n // PROFILE_SENTENCES)
        sample = db.scalars(
            select(zh.c.text)
            .where(func.mod(zh.c.id, step) == 0)
            .limit(PROFILE_SENTENCES)
        ).all()
        script = nlp.detect_script("".join(sample))
        if script:
            meta["script"] = script
    meta["indexed_at"] = datetime.now(UTC).isoformat(timespec="seconds")
    return meta


# ---- export ----

FIXED_COLUMNS = [
    ("hit_id", "命中 ID"),
    ("document_title", "文档"),
    ("document_id", "文档 ID"),
    ("sentence_index", "句序"),
    ("start_cp", "起始码点"),
    ("end_cp", "结束码点"),
    ("left", "左语境"),
    ("match", "命中"),
    ("right", "右语境"),
    ("aligned", "对齐句"),
]


def _rows(items):
    meta_keys = sorted({k for item in items for k in item["metadata"]})
    header = [title for _, title in FIXED_COLUMNS] + meta_keys
    rows = []
    for item in items:
        values = {
            **item,
            "aligned": "\n".join(
                f"[{a['language']}] {a['text']}" for a in item["aligned"]
            ),
        }
        row = [values[key] for key, _ in FIXED_COLUMNS]
        row += [_cell(item["metadata"].get(k)) for k in meta_keys]
        rows.append(row)
    return header, rows


def _cell(value):
    return (
        value if value is None or isinstance(value, str | int | float) else str(value)
    )


def _write_csv(path, header, rows, notice=()):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        if notice:
            writer.writerows([[line] for line in notice] + [[]])
        writer.writerow(header)
        writer.writerows(rows)


def _write_xlsx(path, header, rows, sheet="检索结果", notice=()):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet)
    title_font = Font(name="宋体", bold=True)
    for line in notice:
        ws.append([line])
    if notice:
        ws.append([])
    cells = []
    for title in header:
        cell = WriteOnlyCell(ws, value=title)
        cell.font = title_font
        cells.append(cell)
    ws.append(cells)
    for row in rows:
        ws.append(row)
    wb.save(path)


def run_export(run):
    req = ExportInput.model_validate(run.params)  # seed pinned at enqueue
    items, hits = [], None
    with session() as db:
        user = db.get(User, run.owner_id)
        corpus, restrict = resolve_scope(db, req.scope, user)
        ast, cqp = plan_query(req.query, corpus_tagset(corpus))
        layers = effective.check_layers(db, req.layers, corpus, user)
        with_span_layer(db, cqp, layers, corpus, user)
        notice = source_notice(db, [corpus.id])
        size = (
            10**9 if pure_attr(ast) or query.pure_span(ast) else EXPORT_PAGE
        )  # PostgreSQL answers in one pass
        page = 1
        while True:
            result = search_page(
                db,
                corpus,
                restrict,
                ast,
                cqp,
                layers,
                req,
                page,
                size,
                EXPORT_COUNT_TIMEOUT,
            )
            if hits is None:
                hits = result.counts.hits
            items += enrich(db, result.items)
            run.progress(len(items), hits, "导出")
            if len(result.items) < size:
                break
            page += 1
    header, rows = _rows(items)
    _finish_export(run, "检索结果", header, rows, {"hits": hits}, notice)


SOURCE_FIELDS = (
    ("license", "许可"),
    ("attribution", "署名"),
    ("source_url", "链接"),
    ("version", "版本"),
)


def source_notice(db, corpus_ids):
    """One line per corpus that has a license or attribution (Corpus.meta), for the top of exports."""
    lines = []
    for cid in dict.fromkeys(str(c) for c in corpus_ids if c):
        corpus = db.get(Corpus, uuid.UUID(cid))
        meta = (corpus.meta or {}) if corpus else {}
        if not (meta.get("license") or meta.get("attribution")):
            continue
        parts = [f"来源：{corpus.name}"]
        parts += [f"{label}：{meta[k]}" for k, label in SOURCE_FIELDS if meta.get(k)]
        lines.append("；".join(parts))
    return lines


def _finish_export(run, name, header, rows, extra=None, notice=()):
    """notice: source / license lines written above the table (csv, xlsx) or under the title (docx)."""
    fmt = run.params["format"]
    path = export_path(run.id, fmt)
    path.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "csv":
        _write_csv(path, header, rows, notice)
    elif fmt == "xlsx":
        _write_xlsx(path, header, rows, name, notice)
    else:
        _write_docx(path, name, header, rows, notice)
    filename = f"{name}-{datetime.now():%Y%m%d-%H%M}.{fmt}"
    with transaction() as db:
        run.lock(db)
        run.succeed(db, {"rows": len(rows), **(extra or {}), "filename": filename})


def _write_docx(path, title, header, rows, notice=()):
    """One block per row: the first column as a paragraph, the others as `label：value` lines. 宋体, no fills."""
    doc = DocxDocument()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    heading = doc.add_paragraph().add_run(title)
    heading.bold, heading.font.size = True, Pt(16)
    for line in notice:
        doc.add_paragraph(line)
    for i, row in enumerate(rows, 1):
        doc.add_paragraph(f"{i}. {row[0]}")
        for label, value in zip(header[1:], row[1:]):
            if value not in (None, ""):
                doc.add_paragraph(f"{label}：{value}")
    doc.save(path)


def run_export_stats(run):
    """Rows of a StatResult; the file opens with the method, parameters, scope and token totals (punctuation
    included), then the source / license lines. Decimal values are rounded to 2 places."""
    methods = {
        "frequency": "频次",
        "ngrams": "n-gram",
        "collocations": "搭配",
        "keywords": "关键词",
        "distribution": "分布",
    }
    units = {
        "word": "词形",
        "lemma": "词元",
        "pos": "词性",
        "deprel": "依存关系",
        "span_label": "片段标签",
        "relation_label": "关系标签",
    }
    measures = {
        "logdice": "logDice",
        "mi": "MI",
        "t": "t 值",
        "ll": "LL",
        "log_ratio": "Log Ratio",
    }
    with session() as db:
        stat = db.get(StatResult, uuid.UUID(run.params["stat_id"]))
        if not stat:
            raise JobFailed("STAT_DELETED", "统计结果已删除")
        columns, path = stat.columns, Path(stat.rows_path)
        p, method, totals = stat.params or {}, stat.method, stat.totals or {}
        reference = p.get("reference") or {}
        scope_name, created = stat.scope.get("name", ""), stat.created_at
        index_version = stat.index_version
        notice = source_notice(
            db, [stat.scope.get("corpus_id"), reference.get("corpus_id")]
        )
    label = methods.get(method, method)
    parts = []
    if method == "ngrams":
        parts.append(f"n={p.get('n')}")
    if p.get("unit"):
        parts.append(f"单位 {units.get(p['unit'], p['unit'])}")
    if p.get("node"):
        parts.append(f"节点词 {p['node']}")
    if p.get("window"):
        parts.append(f"窗口 ±{p['window']}")
    if p.get("query"):
        parts.append(f"查询式 {p['query']}")
    if p.get("pos"):
        parts.append(f"词性 {p['pos']}")
    if p.get("pos_pattern"):
        parts.append("词性模式 " + " ".join(t or "*" for t in p["pos_pattern"]))
    if p.get("contains"):
        parts.append(f"包含词 {p['contains']}")
    if method == "collocations" and p.get("measures"):
        parts.append(f"排序指标 {measures.get(p['measures'][0], p['measures'][0])}")
    if method == "keywords" and p.get("measure"):
        parts.append(f"指标 {measures.get(p['measure'], p['measure'])}")
    if method == "distribution":
        field = p.get("field")
        parts.append("按 文档位置" if field == "position" else f"按 {field}")
        if field == "position" and p.get("bins"):
            parts.append(f"区间 {p['bins']}")
    if p.get("min_freq") is not None:
        parts.append(f"最低频次 {p['min_freq']}")
    if p.get("top"):
        parts.append(f"显示前 {p['top']}")
    if p.get("exclude_punct"):
        parts.append("排除标点和数字")
    head = [f"统计：{label}（{'；'.join(parts)}）"]
    head.append(f"范围：{scope_name}，{totals.get('tokens', 0):,} 词（含标点）")
    if reference:
        ref_tokens = (totals.get("reference") or {}).get("tokens", 0)
        head.append(f"参照：{reference.get('name', '')}，{ref_tokens:,} 词（含标点）")
    at = created.astimezone(UTC)
    head.append(f"索引版本：{index_version or '—'}；计算时间：{at:%Y-%m-%d %H:%M} UTC")
    data = json.loads(path.read_text(encoding="utf-8"))["rows"]
    header = [c["label"] for c in columns] + ["查询式"]

    def value(v):
        v = _cell(v)
        return round(v, 2) if isinstance(v, float) else v

    rows = [
        [value(r["values"].get(c["key"])) for c in columns] + [r["search_query"]]
        for r in data
    ]
    _finish_export(run, f"统计-{label}", header, rows, notice=[*head, *notice])


def run_export_evidence(run):
    with session() as db:
        items = db.scalars(
            select(Evidence)
            .where(Evidence.project_id == uuid.UUID(run.params["project_id"]))
            .order_by(Evidence.created_at, Evidence.id)
        ).all()
        # source / license lines for the corpora behind pinned hits and stat rows
        corpus_ids = []
        for e in items:
            if e.ref.get("document_id"):
                doc = db.get(Document, uuid.UUID(e.ref["document_id"]))
                corpus_ids.append(doc and doc.corpus_id)
            if e.ref.get("stat_id"):
                stat = db.get(StatResult, uuid.UUID(e.ref["stat_id"]))
                if stat:
                    reference = (stat.params or {}).get("reference") or {}
                    corpus_ids += [
                        stat.scope.get("corpus_id"),
                        reference.get("corpus_id"),
                    ]
        notice = source_notice(db, corpus_ids)
    # hits pinned with their sentence (services.api.evidence.sentence_evidence): 原文 is that sentence alone,
    # 命中 the matched text, 译文 the aligned sentences; older items and other kinds keep their text
    header = ["原文", "命中", "译文", "出处", "查询式", "备注"]
    rows = []
    for e in items:
        sentence, (a, b) = e.ref.get("sentence"), e.ref.get("match") or (0, 0)
        aligned = "\n".join(
            f"[{x['language']}] {x['text']}" for x in e.ref.get("aligned") or []
        )
        rows.append(
            [
                sentence or e.text,
                sentence[a:b] if sentence else None,
                aligned or None,
                e.source,
                e.ref.get("search_query"),
                e.note,
            ]
        )
    if run.params["format"] != "docx":
        _finish_export(run, "证据篮", header, rows, notice=notice)
        return
    path = export_path(run.id, "docx")
    path.parent.mkdir(parents=True, exist_ok=True)
    _write_evidence_docx(path, items, header, rows, notice)
    with transaction() as db:
        run.lock(db)
        run.succeed(
            db,
            {"rows": len(rows), "filename": f"证据篮-{datetime.now():%Y%m%d-%H%M}.docx"},
        )


def _write_evidence_docx(path, items, header, rows, notice):
    """Like _write_docx, with the hit in bold inside its sentence and no separate 命中 line."""
    doc = DocxDocument()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    heading = doc.add_paragraph().add_run("证据篮")
    heading.bold, heading.font.size = True, Pt(16)
    for line in notice:
        doc.add_paragraph(line)
    for i, (e, row) in enumerate(zip(items, rows), 1):
        p = doc.add_paragraph(f"{i}. ")
        sentence, span = e.ref.get("sentence"), e.ref.get("match")
        if sentence and span:
            a, b = span
            p.add_run(sentence[:a])
            p.add_run(sentence[a:b]).bold = True
            p.add_run(sentence[b:])
        else:
            p.add_run(row[0])
        for label, value in zip(header[2:], row[2:]):
            if value not in (None, ""):
                doc.add_paragraph(f"{label}：{value}")
    doc.save(path)


def _conllu_doc_id(doc, used):
    """`# newdoc id` / sent_id prefix: the document's external id (UD sources) or its title, whitespace as
    `_`; a repeated title gets -2, -3…"""
    base = "_".join((doc.external_id or doc.title or "doc").split()) or "doc"
    out, n = base, 1
    while out in used:
        n += 1
        out = f"{base}-{n}"
    used.add(out)
    return out


def run_export_layer(run):
    """CoNLL-U of a layer's documents with effective values: this layer's confirmed / modified annotations,
    otherwise the system values in `tokens` (for the other kind always the system values). Other formats
    (brat, json, tmx, xlsx): services.worker.annot_tasks."""
    fmt = run.params.get("format", "conllu")
    if fmt != "conllu":
        return annot_tasks.EXPORTS[fmt](run, uuid.UUID(run.params["layer_id"]))
    with session() as db:
        layer = db.get(AnnotationLayer, uuid.UUID(run.params["layer_id"]))
        if not layer:
            raise JobFailed("LAYER_DELETED", "标注层已删除")
        ids = run.params.get("document_ids") or layer.document_ids
        stmt = select(Document).where(Document.corpus_id == layer.corpus_id)
        if ids:
            stmt = stmt.where(Document.id.in_(ids))
        docs = db.scalars(stmt.order_by(Document.created_at, Document.id)).all()
        layers = {"pos": None, "dep": None, layer.kind: effective.eff_layer(layer)}
        name, kind = layer.name, layer.kind
        notice = source_notice(db, [layer.corpus_id])
    parts, sentences, tokens = [f"# {line}\n" for line in notice], 0, 0
    parts.append(f"# 导出口径：{annot_tasks.EXPORT_SCOPE[kind]}\n")
    used = set()
    for i in range(0, len(docs), LOAD_BATCH):
        batch = docs[i : i + LOAD_BATCH]
        run.progress(i, len(docs), "导出")
        with session() as db:
            toks = effective.layer_tokens(db, [d.id for d in batch], layers)
            spans = db.execute(
                select(
                    Sentence.document_id,
                    Sentence.idx,
                    Sentence.start_cp,
                    Sentence.end_cp,
                )
                .where(Sentence.document_id.in_([d.id for d in batch]))
                .order_by(Sentence.document_id, Sentence.idx)
            ).all()
        by_doc = {}
        for doc_id, idx, a, b in spans:
            sent_toks = [em.Token(*t) for t in toks.get(doc_id, {}).get(idx, [])]
            by_doc.setdefault(doc_id, []).append(em.Sentence(a, b, sent_toks))
        out = []
        for d in batch:
            sents = by_doc.get(d.id, [])
            sentences += sum(1 for x in sents if x.tokens)
            tokens += sum(len(x.tokens) for x in sents)
            out.append(
                em.Document(
                    d.title,
                    d.language,
                    Path(d.text_path).read_text(encoding="utf-8"),
                    {},
                    sents,
                    _conllu_doc_id(d, used),
                )
            )
        parts.append(write_conllu(out))
    path = export_path(run.id, "conllu")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(parts), encoding="utf-8", newline="\n")
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in name)
    filename = f"{safe}-{datetime.now():%Y%m%d-%H%M}.conllu"
    with transaction() as db:
        run.lock(db)
        run.succeed(db, {"rows": sentences, "tokens": tokens, "filename": filename})


EXPORTS = {
    "search": run_export,
    "stats": run_export_stats,
    "evidence": run_export_evidence,
    "layer": run_export_layer,
    "kb_terms": kb_tasks.run_export_terms,
    "kb_note": kb_tasks.run_export_note,
}


def run_any_export(run):
    EXPORTS[run.params.get("source", "search")](run)


# ---- stats ----


def _scope_info(db, scope: Scope, corpus):
    sub = db.get(Subcorpus, scope.subcorpus_id) if scope.subcorpus_id else None
    info = {
        "corpus_id": str(corpus.id),
        "subcorpus_id": str(sub.id) if sub else None,
        "name": sub.name if sub else corpus.name,
    }
    if scope.filter is not None:
        info["filter"] = scope.filter
    return info


def _compute(p, target, reference, numeric, positions=None):
    method, top = p["method"], p["top"]
    punct = p.get("exclude_punct", False)  # jobs queued before the option: kept
    if method == "frequency":
        return stats.frequency(
            target, p["unit"], p["pos"], top, p["min_freq"], exclude_punct=punct
        )
    if method == "ngrams":
        return stats.ngrams(
            target,
            p["n"],
            p["unit"],
            p["pos_pattern"],
            top,
            p["min_freq"],
            contains=p.get("contains"),
            exclude_punct=punct,
        )
    if method == "collocations":
        return stats.collocations(
            target,
            p["node"],
            p["window"],
            p["unit"],
            p["pos"],
            p["measures"],
            p["min_freq"],
            top,
            exclude_punct=punct,
        )
    if method == "keywords":
        return stats.keywords(
            target,
            reference,
            p["unit"],
            p["measure"],
            p["min_freq"],
            top,
            exclude_punct=punct,
        )
    return stats.distribution(
        target,
        p["query"],
        p["field"],
        p["bins"],
        numeric=numeric,
        top=top,
        positions=positions,
    )


def _label_frequency(db, p, corpus, scope, user):
    """span_label / relation_label frequency of effective spans / relations in a scope."""
    layers = effective.check_layers(
        db, LayersInput.model_validate(p.get("layers") or {}), corpus, user
    )
    kind = p["unit"].split("_")[0]
    layer = layers[kind]
    if layer is None:
        raise JobFailed("INVALID_LAYER", f"缺少 {kind} 标注层")
    filt = scope_filter(db, scope)
    fields = field_types(db, corpus.id) if filt else {}
    counts, row_count, totals = effective.label_frequency(
        db, corpus, filt, fields, layer, p["unit"], p["top"], p["min_freq"]
    )
    rows = []
    for value, freq, docs in counts:
        per_million = freq / totals["tokens"] * 1e6 if totals["tokens"] else 0.0
        q = (
            f"[{value}]"
            if kind == "span" and _plain(value) and "[" not in value
            else None
        )
        rows.append(
            {
                "key": value,
                "values": {
                    "item": value,
                    "freq": freq,
                    "per_million": per_million,
                    "doc_freq": docs,
                },
                "search_query": q,
            }
        )
    columns = [
        {"key": "item", "label": "标签", "type": "string"},
        {"key": "freq", "label": "频次", "type": "number"},
        {"key": "per_million", "label": "每百万词", "type": "number"},
        {"key": "doc_freq", "label": "文档频率", "type": "number"},
    ]
    return stats.Table(columns, rows, row_count, totals)


def _span_positions(db, p, corpus, scope, user):
    """Token positions of the spans of a pure `[标签]` distribution query (None for other queries)."""
    if not p.get("query"):
        return None
    ast, cqp = plan_query(p["query"], corpus_tagset(corpus))
    spans = query.spans(cqp)
    if not spans or not query.pure_span(ast):
        return None
    layers = effective.check_layers(
        db, LayersInput.model_validate(p.get("layers") or {}), corpus, user
    )
    with_span_layer(db, cqp, layers, corpus, user)
    filt = scope_filter(db, scope)
    fields = field_types(db, corpus.id) if filt else {}
    return effective.span_positions(db, corpus, filt, fields, layers["span"], spans[0])


def _pg_frequency(db, p, corpus, scope, user):
    """Frequency of effective POS tags (with layers.pos) or deprels, counted in PostgreSQL."""
    layers = effective.check_layers(
        db, LayersInput.model_validate(p.get("layers") or {}), corpus, user
    )
    tagset = corpus_tagset(corpus)
    pos_tags = query.resolve_pos(p["pos"], tagset) if p.get("pos") else None
    filt = scope_filter(db, scope)
    fields = field_types(db, corpus.id) if filt else {}
    counts, row_count, totals = effective.frequency(
        db, corpus, filt, fields, p["unit"], layers, pos_tags, p["top"], p["min_freq"]
    )
    rows = []
    for value, freq, docs in counts:
        if p["unit"] == "pos":
            q = (
                f"{{pos={value}}}"
                if _plain(value) and _resolves(value, tagset)
                else None
            )
        else:
            cond = f"pos={p['pos']}," if p.get("pos") else ""
            q = f"{{{cond}dep={value}}}" if _plain(value) else None
        per_million = freq / totals["tokens"] * 1e6 if totals["tokens"] else 0.0
        values = {
            "item": value,
            "freq": freq,
            "per_million": per_million,
            "doc_freq": docs,
        }
        rows.append({"key": value, "values": values, "search_query": q})
    columns = [
        {"key": "item", "label": "项", "type": "string"},
        {"key": "freq", "label": "频次", "type": "number"},
        {"key": "per_million", "label": "每百万词", "type": "number"},
        {"key": "doc_freq", "label": "文档频率", "type": "number"},
    ]
    return stats.Table(columns, rows, row_count, totals)


def _plain(value):
    return not any(c in value for c in ' ,{}|=/"')


def _resolves(value, tagset):
    try:
        return query.resolve_pos(value, tagset) == (value,)
    except query.QueryError:
        return False


def run_stats(run):
    p = run.params
    method = p["method"]
    scope = Scope.model_validate(p["target"] if method == "keywords" else p["scope"])
    params = {k: v for k, v in p.items() if k != "method"}
    reference = numeric = table = positions = None
    pg = method == "frequency" and (
        p["unit"] == "deprel"
        or (p["unit"] == "pos" and (p.get("layers") or {}).get("pos"))
    )
    labels = method == "frequency" and p["unit"] in ("span_label", "relation_label")
    with session() as db:
        user = db.get(User, run.owner_id)
        corpus, restrict = resolve_scope(db, scope, user)
        if labels:
            table = _label_frequency(db, p, corpus, scope, user)
        elif pg:
            table = _pg_frequency(db, p, corpus, scope, user)
        if method == "distribution":
            positions = _span_positions(db, p, corpus, scope, user)
        target = stats.Scope(index_dir(corpus.id), corpus_key(corpus.id), restrict)
        info, base_filter = _scope_info(db, scope, corpus), scope_filter(db, scope)
        index_version = corpus.index_version
        if method == "keywords":
            ref_scope = Scope.model_validate(p["reference"])
            ref_corpus, ref_restrict = resolve_scope(db, ref_scope, user)
            reference = stats.Scope(
                index_dir(ref_corpus.id), corpus_key(ref_corpus.id), ref_restrict
            )
            params["reference"] = _scope_info(db, ref_scope, ref_corpus)
        if method == "distribution":
            numeric = field_types(db, corpus.id).get(p["field"]) == "number"
    run.progress(0, None, "计算")
    try:
        table = table or _compute(p, target, reference, numeric, positions)
    except ValueError as e:  # includes query.QueryError
        raise JobFailed("STATS_INVALID", str(e)) from e
    search_scope = scope.model_dump(mode="json", exclude_none=True)
    for row in table.rows:
        row["search_scope"] = search_scope
        if "bin_value" in row:  # distribution by field: temporary filter for that value
            filt = {**base_filter, p["field"]: row.pop("bin_value")}
            row["search_scope"] = {"corpus_id": str(corpus.id), "filter": filt}
    stat_id = uuid.uuid4()
    path = stats_path(stat_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps({"rows": table.rows}, ensure_ascii=False), encoding="utf-8"
    )
    try:
        with transaction() as db:
            run.lock(db)
            db.add(
                StatResult(
                    id=stat_id,
                    project_id=run.project_id,
                    method=method,
                    params=params,
                    scope=info,
                    index_version=index_version,
                    totals=table.totals,
                    columns=table.columns,
                    rows_path=str(path),
                    row_count=table.row_count,
                    created_by=run.owner_id,
                )
            )
            run.succeed(db, {"stat_id": str(stat_id)})
    except BaseException:
        path.unlink(missing_ok=True)
        raise


# ---- pre-annotation ----


def _layer(db, layer_id, lock=False):
    layer = db.get(AnnotationLayer, layer_id, with_for_update=lock)
    if not layer:
        raise JobFailed("LAYER_DELETED", "标注层已删除")
    return layer


def _document_tokens(db, doc_id):
    return db.execute(
        select(
            Sentence.idx, Token.idx, Token.id, Token.start_cp, Token.end_cp, Token.form
        )
        .add_columns(Token.pos, Token.lemma, Token.head, Token.deprel)
        .join(Token, Token.sentence_id == Sentence.id)
        .where(Sentence.document_id == doc_id)
        .order_by(Sentence.idx, Token.idx)
    ).all()


def _parse_zh(rows):
    """Stanza zh dependencies for a document's stored jieba tokens → {token id: (head, deprel)}."""
    sentences = {}
    for idx, _, tid, a, b, form, pos, lemma, _, _ in rows:
        sentences.setdefault(idx, []).append((tid, em.Token(a, b, form, pos, lemma)))
    doc = em.Document(
        "",
        "zh",
        "",
        {},
        [
            em.Sentence(t[0][1].start, t[-1][1].end, [x for _, x in t])
            for t in sentences.values()
        ],
    )
    nlp.annotate_dependencies_zh(doc)
    return {tid: (t.head, t.deprel) for toks in sentences.values() for tid, t in toks}


def _preannotation(kind, doc_id, language):
    """Pre-annotation rows of one document (read and, for Chinese dependencies, parsed outside any lock)
    and the Stanza values to write back into `tokens` ({token id: (head, deprel)} or None)."""
    with session() as db:
        rows = _document_tokens(db, doc_id)
    source, parsed = "system", None
    if (
        kind == "dep"
        and is_zh(language)
        and any(r.deprel is None or r.head is None for r in rows)
    ):
        parsed, source = _parse_zh(rows), "model:stanza-zh"
    values = []
    for idx, t_idx, tid, _, _, _, pos, _, head, deprel in rows:
        if kind == "pos":
            label, head = pos, None
        else:
            head, label = parsed[tid] if parsed else (head, deprel)
        values.append(
            {
                "document_id": doc_id,
                "sentence_index": idx,
                "token_index": t_idx + 1,
                "label": label,
                "head": head,
                "model_label": label,
                "model_head": head,
                "source": source,
                "status": "suggested",
            }
        )
    return values, parsed


def _write_preannotation(db, layer_id, values, parsed, annotators):
    conn = db.connection()
    if parsed:  # fill the system layer where it has no dependencies yet
        conn.execute(
            update(Token)
            .where(Token.id == bindparam("tid"))
            .values(
                head=func.coalesce(Token.head, bindparam("h")),
                deprel=func.coalesce(Token.deprel, bindparam("d")),
            ),
            [{"tid": tid, "h": h, "d": d} for tid, (h, d) in parsed.items()],
        )
    if not values:
        return
    stmt = pg_insert(Annotation)
    conn.execute(
        stmt.on_conflict_do_update(
            constraint="uq_annotations_token",
            set_={
                "label": stmt.excluded.label,
                "head": stmt.excluded.head,
                "model_label": stmt.excluded.model_label,
                "model_head": stmt.excluded.model_head,
                "source": stmt.excluded.source,
                "revision": Annotation.revision + 1,
                "updated_at": func.now(),
            },
            where=Annotation.status == "suggested",
        ),
        [
            {"layer_id": layer_id, "annotator_id": a, **v}
            for a in annotators
            for v in values
        ],
    )


def run_preannotate(run):
    """Fill a layer from `tokens` (contracts/p3.md §三): pos = tokens.pos; dep = tokens.head / deprel, or for
    Chinese documents without dependencies Stanza zh over the stored jieba tokens (also written back to
    `tokens`). Re-runs only overwrite tokens still `suggested`. One transaction per document; progress counts
    documents."""
    layer_id = uuid.UUID(run.params["layer_id"])
    with session() as db:
        token_kind = _layer(db, layer_id).kind in ("pos", "dep")
    if not token_kind:
        return annot_tasks.run_preannotate_p5(run)
    with transaction() as db:
        layer = _layer(db, layer_id, lock=True)
        layer.status, layer.error, layer.job_id = "preannotating", None, run.id
        kind = layer.kind
        annotators = layer_annotators(layer)
        ids = run.params.get("document_ids") or layer.document_ids
        stmt = select(Document.id, Document.language).where(
            Document.corpus_id == layer.corpus_id
        )
        if ids:
            stmt = stmt.where(Document.id.in_(ids))
        docs = db.execute(stmt.order_by(Document.created_at, Document.id)).all()
    try:
        for i, (doc_id, language) in enumerate(docs):
            run.progress(i, len(docs), "预标")
            values, parsed = _preannotation(kind, doc_id, language)
            with transaction() as db:
                layer = _layer(db, layer_id, lock=True)
                _write_preannotation(db, layer_id, values, parsed, annotators)
                refresh_progress(db, layer)
        with transaction() as db:
            run.lock(db)
            layer = _layer(db, layer_id, lock=True)
            refresh_progress(db, layer)
            layer.status = "ready"
            run.succeed(
                db,
                {
                    "layer_id": str(layer_id),
                    "documents": len(docs),
                    "tokens": layer.progress["total"],
                },
            )
    except BaseException as e:
        message = (
            "预标已取消"
            if isinstance(e, Cancelled)
            else getattr(e, "message", str(e))[:500]
        )
        with transaction() as db:
            layer = db.get(AnnotationLayer, layer_id, with_for_update=True)
            if layer:
                layer.status, layer.error = "failed", message
        raise


HANDLERS = {
    "import": run_import,
    "index": run_index,
    "export": run_any_export,
    "stats": run_stats,
    "preannotate": run_preannotate,
    "kb_import": kb_tasks.run_kb_import,
    "kb_summarize": kb_tasks.run_kb_summarize,
    "kb_extract": kb_tasks.run_kb_extract,
}
